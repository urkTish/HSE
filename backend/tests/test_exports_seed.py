"""Exports (AC33, rule 49) and seed data (AC34)."""

import csv
import io
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import AuditAction
from app.models import AuditEntry, Contractor, User
from app.seed import seed
from tests.conftest import PASSWORD, Api, Ids


def _csv(res: object) -> list[list[str]]:
    text = res.content.decode("utf-8-sig")  # type: ignore[attr-defined]
    return list(csv.reader(io.StringIO(text)))


def test_AC33_viewer_export_has_no_contacts(api: Api, db: Session) -> None:
    res = api.as_("sarah.mitchell").get("/api/v1/exports/contractors", params={"format": "csv"})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    rows = _csv(res)
    header = rows[0]
    assert not {"primary_contact_name", "primary_contact_mobile", "primary_contact_email"} & set(
        header
    )
    assert len(rows) - 1 == 6
    entry = db.scalars(select(AuditEntry).where(AuditEntry.action == AuditAction.export)).one()
    assert entry.details is not None and entry.details["row_count"] == 6
    assert entry.details["dataset"] == "contractors"


def test_rule49_officer_export_includes_contacts_xlsx(api: Api) -> None:
    res = api.as_("noura.qahtani").get("/api/v1/exports/contractors", params={"format": "xlsx"})
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "primary_contact_email" in header


def test_export_scoping(api: Api, ids: Ids) -> None:
    omar = api.as_("omar.siddiqui")
    rows = _csv(omar.get("/api/v1/exports/zones", params={"project_id": ids.project("ANIA-EXP")}))
    assert {r[1] for r in rows[1:]} == {"Z-APR-21", "Z-TWB", "Z-ILS33R"}
    # permit issuer has no export capability
    khalid = api.as_("khalid.otaibi")
    assert khalid.get("/api/v1/exports/contractors").status_code == 403
    # project_id is required for project-scoped lists
    assert omar.get("/api/v1/exports/zones").status_code == 422
    rep = api.as_("ahmed.zahrani")
    eng = _csv(
        rep.get("/api/v1/exports/engagements", params={"project_id": ids.project("ANIA-EXP")})
    )
    assert {r[0] for r in eng[1:]} == {"RAWABI", "NAJD", "GULFPAVE", "SAHARA"}
    users = _csv(api.as_("faisal.harbi").get("/api/v1/exports/users"))
    assert "email" in users[0] and len(users) == 9


def test_AC34_seed_data_is_fake(db: Session) -> None:
    mobile = re.compile(r"^\+96650000\d{4}$")
    users = db.scalars(select(User)).all()
    assert len(users) == 8
    for u in users:
        assert u.email.endswith("@example.com")
        assert u.mobile and mobile.match(u.mobile)
    for c in db.scalars(select(Contractor)).all():
        assert c.primary_contact_email.endswith("@example.com")
        assert mobile.match(c.primary_contact_mobile)


def test_seed_is_idempotent(db: Session) -> None:
    before = len(db.scalars(select(User)).all())
    seed(db, PASSWORD)
    db.commit()
    assert len(db.scalars(select(User)).all()) == before
