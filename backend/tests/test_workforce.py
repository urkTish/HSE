"""Workforce returns, month lock and the CSV/Excel import (spec 1-dashboard §3.2, §4.1,
AC2-AC12)."""

import io
from datetime import date, timedelta
from typing import Any

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.core.clock import now
from app.core.config import get_settings
from app.core.enums import UserStatus
from app.core.security import hash_password
from app.db.session import get_sessionmaker
from app.models import User, WorkforceImportBatch, WorkforceReturn
from app.services.workforce_import import COLUMNS
from tests.conftest import PASSWORD, Api, Ids

API = "/api/v1"
HEADER = [
    "work_date",
    "project_code",
    "site_code",
    "zone_code",
    "contractor_code",
    "shift",
    "headcount",
    "man_hours",
    "no_work",
]


def row(d: str, site: str, eng: str, shift: str, hc: int, mh: int, **kw: Any) -> list[str]:
    return [
        d,
        kw.get("project", "ANIA-EXP"),
        site,
        kw.get("zone", ""),
        eng,
        shift,
        str(hc),
        str(mh),
        kw.get("no_work", ""),
    ]


def csv(rows: list[list[str]], header: list[str] = HEADER) -> bytes:
    return "\n".join(",".join(r) for r in [header, *rows]).encode()


def upload(
    c: TestClient, pid: str, content: bytes, mode: str = "insert_only", name: str = "f.csv"
) -> Any:
    res = c.post(
        f"{API}/projects/{pid}/workforce-imports",
        files={"file": (name, content)},
        data={"mode": mode},
    )
    assert res.status_code == 201, res.text
    return res.json()


def codes_of(batch: dict[str, Any], row_no: int) -> list[str]:
    return next(r["codes"] for r in batch["report"] if r["row_no"] == row_no)


def ret_body(ids: Ids, **over: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "site_id": ids.site("S-AIR"),
        "engagement_id": ids.engagement("GULFPAVE"),
        "work_date": "2026-09-14",
        "shift": "night",
        "headcount": 10,
        "man_hours": "100",
    }
    body.update(over)
    return body


def returns(db: Session) -> int:
    db.expire_all()
    return int(db.scalar(select(func.count()).select_from(WorkforceReturn)) or 0)


def test_AC2_duplicate_return_rejected(api: Api, ids: Ids) -> None:
    c = api.as_("noura.qahtani")
    url = f"{API}/projects/{ids.project('ANIA-EXP')}/workforce-returns"
    assert c.post(url, json=ret_body(ids)).status_code == 201
    res = c.post(url, json=ret_body(ids, headcount=12, man_hours="120"))
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "DUPLICATE_RETURN"
    # a different shift is a different key
    assert c.post(url, json=ret_body(ids, shift="day")).status_code == 201


def _hundred_valid() -> list[list[str]]:
    rows = []
    for i in range(25):
        d = (date(2026, 9, 1) + timedelta(days=i)).isoformat()
        for site in ("S-AIR", "S-LAND"):
            for shift in ("day", "night"):
                rows.append(row(d, site, "RAWABI", shift, 20, 200))
    return rows


def test_AC3_dry_run_reports_E08_writes_nothing(api: Api, ids: Ids, db: Session) -> None:
    c = api.as_("noura.qahtani")
    rows = [*_hundred_valid(), row("2026-09-14", "S-AIR", "GULFPAVE", "night", 10, 200)]
    b = upload(c, ids.project("ANIA-EXP"), csv(rows))
    assert b["counts"]["rows_total"] == 101
    assert b["counts"]["rows_error"] == 1
    assert [x for x in codes_of(b, 101) if x.startswith("E")] == ["E08"]
    assert returns(db) == 0
    res = c.post(f"{API}/workforce-imports/{b['id']}/commit")
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "IMPORT_HAS_ERRORS"
    assert returns(db) == 0


def test_AC4_commit_with_warnings(api: Api, ids: Ids, db: Session) -> None:
    c = api.as_("noura.qahtani")
    rows = [
        row("2026-09-20", "S-AIR", "RAWABI", "day", 10, 130),  # 13 h/person > 12 → W01
        row("2026-09-20", "S-LAND", "RAWABI", "day", 10, 125),  # W01
        row("2026-09-20", "S-LAND", "NAJD", "day", 10, 100),
    ]
    b = upload(c, ids.project("ANIA-EXP"), csv(rows))
    assert b["counts"]["rows_error"] == 0
    assert b["counts"]["rows_warning"] == 2
    res = c.post(f"{API}/workforce-imports/{b['id']}/commit")
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "committed"
    assert res.json()["counts"]["rows_inserted"] == 3
    db.expire_all()
    rs = list(db.scalars(select(WorkforceReturn)))
    assert len(rs) == 3
    assert {r.status.value for r in rs} == {"submitted"}
    assert {r.source.value for r in rs} == {"import"}
    warned = [r for r in rs if any(w["code"] == "W01" for w in r.warnings or [])]
    assert len(warned) == 2


def test_AC5_expired_dry_run(api: Api, ids: Ids, db: Session) -> None:
    c = api.as_("noura.qahtani")
    b = upload(
        c, ids.project("ANIA-EXP"), csv([row("2026-09-20", "S-AIR", "RAWABI", "day", 5, 40)])
    )
    db.execute(
        update(WorkforceImportBatch)
        .where(WorkforceImportBatch.id == b["id"])
        .values(
            created_at=now() - timedelta(minutes=61),
            expires_at=now() - timedelta(minutes=1),
        )
    )
    db.commit()
    res = c.post(f"{API}/workforce-imports/{b['id']}/commit")
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "IMPORT_EXPIRED"
    assert returns(db) == 0


def test_AC6_insert_only_E10_and_upsert_replaces_verified(api: Api, ids: Ids, db: Session) -> None:
    noura, mgr = api.as_("noura.qahtani"), api.as_("faisal.harbi")
    pid = ids.project("ANIA-EXP")
    r = noura.post(
        f"{API}/projects/{pid}/workforce-returns", json=ret_body(ids, submit=True)
    ).json()
    v = mgr.post(f"{API}/workforce-returns/{r['id']}/transitions", json={"to_status": "verified"})
    assert v.status_code == 200, v.text
    same = csv([row("2026-09-14", "S-AIR", "GULFPAVE", "night", 12, 120)])
    b = upload(noura, pid, same)
    assert codes_of(b, 1) == ["E10"]
    b2 = upload(noura, pid, same, mode="upsert")
    assert b2["counts"]["rows_error"] == 0
    res = noura.post(f"{API}/workforce-imports/{b2['id']}/commit")
    assert res.status_code == 200, res.text
    assert res.json()["counts"]["rows_replaced"] == 1
    db.expire_all()
    rows = list(db.scalars(select(WorkforceReturn)))
    assert len(rows) == 1
    assert rows[0].status.value == "submitted"
    assert rows[0].headcount == 12
    hist = mgr.get(f"{API}/history/workforce_return/{rows[0].id}").json()["items"]
    changed = [h for h in hist if h["before"] and h["after"]]
    assert any(
        h["before"].get("headcount") == 10 and h["after"].get("headcount") == 12 for h in changed
    )


def _lock_august(api: Api, ids: Ids) -> None:
    res = api.as_("faisal.harbi").post(
        f"{API}/projects/{ids.project('ANIA-EXP')}/workforce-months/2026-08/lock", json={}
    )
    assert res.status_code == 200, res.text


def test_AC7_locked_month_E11_in_both_modes(api: Api, ids: Ids) -> None:
    _lock_august(api, ids)
    c = api.as_("noura.qahtani")
    content = csv([row("2026-08-20", "S-AIR", "RAWABI", "day", 10, 100)])
    for mode in ("insert_only", "upsert"):
        b = upload(c, ids.project("ANIA-EXP"), content, mode=mode)
        assert "E11" in codes_of(b, 1), mode


def test_AC8_contractor_rep_scope(api: Api, ids: Ids) -> None:
    c = api.as_("ahmed.zahrani")
    rows = [
        row("2026-09-20", "S-TWR", "QIMMA", "day", 10, 100),
        row("2026-09-20", "S-TWR", "QIMMA", "day", 10, 100, project="RBT-52"),
        row("2026-09-20", "S-LAND", "SAHARA", "day", 10, 100),
    ]
    b = upload(c, ids.project("ANIA-EXP"), csv(rows))
    assert {"E12", "E04"} <= set(codes_of(b, 1))
    assert {"E01", "E04", "E12"} <= set(codes_of(b, 2))
    report = {r["row_no"]: r for r in b["report"]}
    assert 3 not in report or not any(x.startswith("E") for x in report[3]["codes"])
    assert b["counts"]["rows_error"] == 2


def test_AC9_arabic_excel_parses_like_english(api: Api, ids: Ids) -> None:
    c = api.as_("noura.qahtani")
    pid = ids.project("ANIA-EXP")
    ar = c.get(f"{API}/workforce-imports/template", params={"format": "xlsx", "headers": "ar"})
    assert ar.status_code == 200
    ar_header = [
        x.value
        for x in next(openpyxl.load_workbook(io.BytesIO(ar.content)).active.iter_rows(max_row=1))
    ]
    assert "تاريخ_العمل" in ar_header and "رمز_المقاول" in ar_header
    key = dict(zip(HEADER, row("2026-09-20", "S-AIR", "GULFPAVE", "night", 10, 160), strict=True))
    key_ar = {**key, "shift": "ليلية"}
    ar_of = {k: a for k, a, _ in COLUMNS}
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append([ar_of[k] for k in HEADER])
    ws.append([key_ar[k] for k in HEADER])
    buf = io.BytesIO()
    wb.save(buf)
    b_ar = upload(c, pid, buf.getvalue(), name="ar.xlsx")
    b_en = upload(c, pid, csv([[key[k] for k in HEADER]]))
    assert b_ar["counts"] == b_en["counts"]
    get = {
        b["id"]: c.get(
            f"{API}/workforce-imports/{b['id']}", params={"include_ok_rows": True}
        ).json()
        for b in (b_ar, b_en)
    }
    r_ar, r_en = (get[b["id"]]["report"][0] for b in (b_ar, b_en))
    for k in ("work_date", "site_code", "contractor_code", "shift", "codes", "status"):
        assert r_ar[k] == r_en[k], k
    assert r_ar["shift"] == "night"


def test_AC10_no_work_with_headcount_E13(api: Api, ids: Ids) -> None:
    c = api.as_("noura.qahtani")
    b = upload(
        c,
        ids.project("ANIA-EXP"),
        csv([row("2026-09-20", "S-AIR", "RAWABI", "day", 5, 0, no_work="Y")]),
    )
    assert "E13" in codes_of(b, 1)


def test_AC11_suspended_contractor(api: Api, ids: Ids) -> None:
    mgr = api.as_("faisal.harbi")
    pid = ids.project("RBT-52")
    rows = [row("2026-08-10", "S-TWR", "DLIFT", "day", 8, 64, project="RBT-52")]
    b = upload(mgr, pid, csv(rows))
    assert b["counts"]["rows_error"] == 0
    assert "W03" in codes_of(b, 1)
    invite = mgr.post(
        f"{API}/users",
        json={
            "email": "dlift.rep@example.com",
            "full_name_en": "Saeed Al-Dosari",
            "employer_type": "contractor",
            "employer_contractor_id": ids.contractor("DLIFT"),
            "role_assignments": [
                {
                    "role": "contractor_hse_rep",
                    "project_id": pid,
                    "contractor_engagement_id": ids.engagement("DLIFT"),
                }
            ],
        },
    )
    assert invite.status_code == 201, invite.text
    with get_sessionmaker()() as s:
        s.execute(
            update(User)
            .where(User.email == "dlift.rep@example.com")
            .values(
                status=UserStatus.active,
                password_hash=hash_password(PASSWORD),
                privacy_notice_version=get_settings().privacy_notice_version,
            )
        )
        s.commit()
    res = api.as_("dlift.rep").post(
        f"{API}/projects/{pid}/workforce-returns",
        json=ret_body(
            ids,
            site_id=ids.site("S-TWR"),
            engagement_id=ids.engagement("DLIFT"),
            work_date="2026-08-11",
        ),
    )
    assert res.status_code == 403
    assert res.json()["detail"]["code"] == "CONTRACTOR_SUSPENDED"


def test_AC12_locked_row_correction_restates_month(api: Api, ids: Ids) -> None:
    noura, mgr = api.as_("noura.qahtani"), api.as_("faisal.harbi")
    pid = ids.project("ANIA-EXP")
    r = noura.post(
        f"{API}/projects/{pid}/workforce-returns",
        json=ret_body(ids, work_date="2026-08-14", submit=True),
    ).json()
    assert (
        mgr.post(
            f"{API}/workforce-returns/{r['id']}/transitions", json={"to_status": "verified"}
        ).status_code
        == 200
    )
    _lock_august(api, ids)
    assert (
        noura.patch(f"{API}/workforce-returns/{r['id']}", json={"headcount": 11}).status_code == 409
    )
    un = mgr.post(
        f"{API}/projects/{pid}/workforce-months/2026-08/unlock",
        json={"reason": "GULFPAVE night hours under-reported"},
    )
    assert un.status_code == 200, un.text
    back = mgr.post(
        f"{API}/workforce-returns/{r['id']}/transitions",
        json={"to_status": "submitted", "reason": "Correct night hours"},
    )
    assert back.status_code == 200, back.text
    fix = mgr.patch(f"{API}/workforce-returns/{r['id']}", json={"man_hours": "120"})
    assert fix.status_code == 200, fix.text
    q = {"project_id": pid, "metric": "K-01", "period": "month", "anchor": "2026-08-01"}
    ctx = mgr.get(f"{API}/kpi/metrics", params=q).json()["context"]
    assert "2026-08" in ctx["restated_months"]
    months = mgr.get(f"{API}/projects/{pid}/workforce-months").json()["items"]
    assert next(m for m in months if m["month"] == "2026-08")["restated"] is True
    hist = mgr.get(f"{API}/history/workforce_return/{r['id']}").json()["items"]
    assert any(
        (h["before"] or {}).get("man_hours") == "100.00" or (h["before"] or {}).get("man_hours")
        for h in hist
        if h["after"] and "man_hours" in h["after"]
    )
