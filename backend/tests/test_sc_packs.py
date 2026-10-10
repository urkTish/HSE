"""Phase 6g report packs and distribution (6g §9 AC 31-42)."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.scorecard_enums import (
    RpAction,
    RpDeliveryStatus,
    RpMemberKind,
    RpStatus,
    RpType,
    XpPurpose,
)
from app.models import Attachment, RpDelivery, RpPack
from app.schemas.scorecard import (
    RpDistributionUpdate,
    RpMemberInput,
    RpPackCreate,
    RpPackTransition,
    RpPackUpdate,
    ScSettingsUpdate,
)
from app.services.scorecard import config, distribution, packs, render
from tests.conftest import Api
from tests.sc_helpers import API, P, expect, notified, pack, project, tick, uid

pytestmark = pytest.mark.usefixtures("sc_seed", "sc_clock")


def _bytes(db: Session, pk: RpPack, kind: str) -> bytes:
    a = db.get(Attachment, uuid.UUID(pk.files[kind]["attachment_id"]))
    assert a is not None
    return (Path(get_settings().storage_dir) / a.storage_key).read_bytes()


def _texts(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    seen: list[str] = []
    orig_cell, orig_multi = render._Pdf.cell, render._Pdf.multi_cell

    def cell(self: Any, *a: Any, **k: Any) -> Any:
        seen.append(str(k.get("text", a[2] if len(a) > 2 else "")))
        return orig_cell(self, *a, **k)

    def multi(self: Any, *a: Any, **k: Any) -> Any:
        seen.append(str(k.get("text", a[2] if len(a) > 2 else "")))
        return orig_multi(self, *a, **k)

    monkeypatch.setattr(render._Pdf, "cell", cell)
    monkeypatch.setattr(render._Pdf, "multi_cell", multi)
    return seen


def _go(db: Session, pk: RpPack, prep: str, reviewer: str, issuer: str, **kw: Any) -> Any:
    packs.transition(db, P(db, prep), pk.id, RpPackTransition(action=RpAction.submit_for_review))
    packs.transition(db, P(db, reviewer), pk.id, RpPackTransition(action=RpAction.review))
    return lambda: packs.transition(
        db, P(db, issuer), pk.id, RpPackTransition(action=RpAction.issue, **kw)
    )


def test_mcr_sg5_revisions_ac31_35_36(
    db: Session, api: Api, monkeypatch: pytest.MonkeyPatch
) -> None:
    r0, r1 = pack(db, "MCR-ANIA-EXP-2026-08", 0), pack(db, "MCR-ANIA-EXP-2026-08", 1)
    assert (
        r0.status == RpStatus.superseded
        and r0.superseded_by_revision == 1
        and r0.revised_since_issue
    )
    assert r1.status == RpStatus.issued and (r1.reissue_reason or "").startswith("August restated")
    seen = _texts(monkeypatch)
    render.pdf(packs.to_doc(db, r0), "en", 1)
    assert "SUPERSEDED BY Rev 1" in seen
    assert any(
        s.startswith("MCR-ANIA-EXP-2026-08 · Rev 0 · page 1 of ")
        and r0.snapshot_hash
        and r0.snapshot_hash[:12] in s
        for s in seen
    )
    for kind in ("pdf_en", "pdf_ar"):
        assert _bytes(db, r1, kind).startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(_bytes(db, r1, "xlsx")))
    ws = wb.worksheets[0]
    heads = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=6)]
    assert any(
        any(isinstance(v, str) and any("؀" <= ch <= "ۿ" for ch in v) for v in row) for row in heads
    )
    assert any(
        isinstance(c.value, int | float)
        for w in wb.worksheets
        for row in w.iter_rows()
        for c in row
    )
    expect(
        "PACK_ISSUED_IMMUTABLE",
        lambda: packs.update(
            db, P(db, "faisal.harbi"), r1.id, RpPackUpdate(review_comment="late edit")
        ),
    )
    r = api.as_("faisal.harbi").get(
        f"{API}/kpi/scorecards",
        params={
            "project_id": str(r1.project_id),
            "period": "custom",
            "start": "2026-09-01",
            "end": "2026-09-30",
        },
    )
    k135 = next(m for m in r.json()["metrics"] if m["metric"] in ("K135", "K-135"))
    assert k135["display"] == "100.0 %"
    assert (
        packs.restatement(db, tick(2026, 10, 13, 7)) == 0
    )  # frozen packs, alert only once (seeded)


def test_mcr_september_ac32_33(db: Session) -> None:
    pk = pack(db, "MCR-ANIA-EXP-2026-09")
    assert pk.status == RpStatus.draft
    keys = [s["key"] for s in pk.snapshot["sections"]]
    assert keys[0] == "p1:executive_summary" and keys[-1] == "appendix" and "module:heat" in keys
    assert "module:followup" not in keys and "Incident follow-up" in pk.snapshot["not_live"]
    # design pass: "Contractor scorecards" once (RP-1 (3)); section 7 is the ranking summary
    titles = [s["title_en"] for s in pk.snapshot["sections"]]
    assert titles.count("Contractor scorecards") == 1, titles
    s7 = next((s for s in pk.snapshot["sections"] if s["key"] == "p1:contractor_performance"), None)
    if s7 is not None:
        assert [c[0] for c in s7["tables"][0]["columns"]] == list(packs.SUMMARY_COLS)
    issue = _go(db, pk, "noura.qahtani", "noura.qahtani", "faisal.harbi")
    expect("SCORECARDS_NOT_FINAL", issue)
    packs.transition(
        db,
        P(db, "faisal.harbi"),
        pk.id,
        RpPackTransition(
            action=RpAction.issue,
            scorecards_provisional=True,
            provisional_reason="Client asked for the report before finalisation.",
        ),
    )
    assert pk.status == RpStatus.issued and pk.scorecards_provisional
    sc = next(s for s in pk.snapshot["sections"] if s["key"] == "scorecards")
    assert (sc.get("watermark") or "").lower().startswith("provisional")
    assert [s["title_en"] for s in pk.snapshot["sections"]].count("Contractor scorecards") == 1
    assert set(pk.files) >= {"pdf_en", "pdf_ar", "xlsx"}


def test_mcr_rbt_checks_ac34_42(db: Session) -> None:
    t = tick(2026, 10, 12, 7)
    rbt = project(db, "RBT-52")
    out = packs.run_daily(db, t)
    assert out is not None
    assert db.scalar(select(RpPack).where(RpPack.doc_no == "MCR-RBT-52-2026-09")) is None
    due = notified(db, "report_pack_due")
    assert "lina.haddad" in due and "faisal.harbi" in due
    d = packs.create(
        db,
        P(db, "lina.haddad"),
        RpPackCreate(report_type=RpType.MCR, project_id=rbt.id, month="2026-09"),
    )
    pk = db.get(RpPack, d.id)
    assert pk is not None
    expect("SOURCE_REPORT_NOT_PUBLISHED", _go(db, pk, "lina.haddad", "lina.haddad", "faisal.harbi"))
    packs.transition(db, P(db, "faisal.harbi"), pk.id, RpPackTransition(action=RpAction.review))
    expect(
        "SELF_REVIEW",
        lambda: packs.transition(
            db, P(db, "faisal.harbi"), pk.id, RpPackTransition(action=RpAction.issue)
        ),
    )


def test_distribution_ac37_38(db: Session) -> None:
    pid = project(db, "ANIA-EXP").id
    f = P(db, "faisal.harbi")
    ahmed = RpMemberInput(kind=RpMemberKind.user, user_id=uid(db, "ahmed.zahrani"))
    expect(
        "RECIPIENT_SCOPE",
        lambda: distribution.put_list(
            db, f, pid, RpType.MCR, RpDistributionUpdate(members=[ahmed])
        ),
    )
    ext = RpMemberInput(
        kind=RpMemberKind.external,
        email="pmc@example.com",
        display_name_en="PMC",
        acknowledge_disclosure=True,
    )
    config.update_settings(db, f, pid, ScSettingsUpdate(external_distribution_enabled=False))
    expect(
        "EXTERNAL_NOT_ALLOWED",
        lambda: distribution.put_list(db, f, pid, RpType.MCR, RpDistributionUpdate(members=[ext])),
    )
    config.update_settings(
        db,
        f,
        pid,
        ScSettingsUpdate(external_distribution_enabled=True, external_domains=["example.com"]),
    )
    sarah = RpMemberInput(kind=RpMemberKind.user, user_id=uid(db, "sarah.mitchell"))
    saved = distribution.put_list(
        db, f, pid, RpType.MCR, RpDistributionUpdate(members=[sarah, ext])
    )
    assert len(saved.members) == 2
    r1 = pack(db, "MCR-ANIA-EXP-2026-08", 1)
    t0 = tick(2026, 10, 12, 8)
    distribution.distribute(db, r1)
    rows = db.scalars(
        select(RpDelivery).where(RpDelivery.pack_id == r1.id, RpDelivery.sent_at >= t0)
    ).all()
    assert {r.user_id or r.email for r in rows} == {uid(db, "sarah.mitchell"), "pmc@example.com"}
    ext_row = next(r for r in rows if r.email == "pmc@example.com")
    assert any(a.endswith(".xlsx") for a in ext_row.attachments) and any(
        a.endswith(".pdf") for a in ext_row.attachments
    )
    assert "sarah.mitchell" in notified(db, "report_pack_issued", t0)
    t = tick(2026, 10, 12, 9)
    distribution.bounce(db, ext_row.id, "550 mailbox unavailable")
    assert ext_row.status == RpDeliveryStatus.bounced
    assert notified(db, "delivery_bounced", t).get("faisal.harbi")


def test_osha_sg6_names_ac39_40(db: Session) -> None:
    pid = project(db, "ANIA-EXP").id
    f = P(db, "faisal.harbi")
    d = packs.create(
        db,
        f,
        RpPackCreate(
            report_type=RpType.OSHA300, project_id=pid, year=2026, period_end=date(2026, 9, 30)
        ),
    )
    pk = db.get(RpPack, d.id)
    assert pk is not None
    secs = {s["key"]: s for s in pk.snapshot["sections"]}
    summ = {r["item"]: str(r["value"]) for r in secs["summary"]["tables"][0]["rows"]}
    assert [
        summ[k]
        for k in (
            "Deaths",
            "Days-away cases",
            "Job transfer or restriction",
            "Other recordable",
            "Total recordable",
            "Days away",
        )
    ] == ["0", "2", "7", "15", "24", "35"]
    assert summ["Man-hours (K-01)"].replace(",", "") == "7030000" and summ["TRIR (K-21)"] == "0.68"
    log = secs["log"]["tables"][0]["rows"]
    assert all(r["person"].startswith("Person ") or r["person"] == "Privacy case" for r in log)
    expect(
        "PURPOSE_REQUIRED",
        lambda: packs.create(
            db,
            f,
            RpPackCreate(
                report_type=RpType.OSHA300,
                project_id=pid,
                year=2026,
                period_end=date(2026, 9, 30),
                with_names=True,
            ),
        ),
    )
    db.delete(pk)
    db.flush()
    n = packs.create(
        db,
        f,
        RpPackCreate(
            report_type=RpType.OSHA300,
            project_id=pid,
            year=2026,
            period_end=date(2026, 9, 30),
            with_names=True,
            purpose=XpPurpose.legal,
            purpose_text="Insurer claim review.",
        ),
    )
    named = db.get(RpPack, n.id)
    assert named is not None and named.with_names
    privacy = [
        r
        for r in named.snapshot["sections"][0]["tables"][0]["rows"]
        if r["person"] == "Privacy case"
    ]
    assert all(r["person"] == "Privacy case" for r in privacy)
    distribution.put_list(
        db,
        f,
        pid,
        RpType.OSHA300,
        RpDistributionUpdate(
            members=[RpMemberInput(kind=RpMemberKind.user, user_id=uid(db, "sarah.mitchell"))]
        ),
    )
    expect("RECIPIENT_SCOPE", _go(db, named, "noura.qahtani", "noura.qahtani", "faisal.harbi"))
    distribution.put_list(db, f, pid, RpType.OSHA300, RpDistributionUpdate(members=[]))
    packs.transition(db, f, named.id, RpPackTransition(action=RpAction.issue))
    exp = datetime.fromisoformat(named.files["pdf_en"]["expires_at"])
    assert (
        timedelta(hours=23)
        < exp - datetime(2026, 10, 12, 7, tzinfo=UTC)
        <= timedelta(hours=24, minutes=1)
    )


def test_scp_ac41(db: Session) -> None:
    for short in ("NAJD", "SAHARA"):
        pk = pack(db, f"SCP-ANIA-EXP-{short}-2026-08")
        text = str(pk.snapshot)
        assert all(o not in text for o in ("RAWABI", "GULFPAVE", "NAJD", "SAHARA") if o != short)
        summ = {r["item"]: r["value"] for r in pk.snapshot["sections"][0]["tables"][0]["rows"]}
        assert summ["Rank"].endswith(" of 4") and summ["Project median"]
        rows = db.scalars(
            select(RpDelivery).where(
                RpDelivery.pack_id == pk.id, RpDelivery.user_id == uid(db, "tariq.mutairi")
            )
        ).all()
        assert rows
