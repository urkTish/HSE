"""The 6g renderer (spec 6g RP-5, EX-6, EX-11): A4 PDFs in English (LTR) and Arabic (RTL, shaped
with HarfBuzz, embedded DejaVu Sans; numbers, codes and refs stay left-to-right through the
bidirectional algorithm) and XLSX workbooks with EN and AR header rows and numeric cells.

Every PDF page carries the footer `doc_no · Rev n · page x of y · <hash12> · CONFIDENTIAL` in a
core Latin font. fpdf2 + uharfbuzz were added for this (DECISIONS D-222)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.comments import Comment

FONTS = Path(__file__).resolve().parents[2] / "data" / "fonts"
PROVISIONAL = "Provisional — subject to contractor comment / مبدئي — قابل لملاحظات المقاول"


@dataclass
class Table:
    columns: list[tuple[str, str, str]]  # key, EN header, AR header
    rows: list[dict[str, Any]]
    title_en: str = ""
    title_ar: str = ""
    sheet: str | None = None  # XLSX sheet name (EX: one sheet per KPI table)


@dataclass
class Section:
    key: str
    title_en: str
    title_ar: str
    tables: list[Table] = field(default_factory=list)
    paragraphs_en: list[str] = field(default_factory=list)
    paragraphs_ar: list[str] = field(default_factory=list)
    watermark: str | None = None


@dataclass
class Doc:
    doc_no: str
    revision: int
    title_en: str
    title_ar: str
    snapshot_hash: str
    control: list[tuple[str, str, str]]  # EN label, AR label, value
    sections: list[Section]
    footer_note: str = "CONFIDENTIAL"


def _txt(v: Any) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, Decimal):
        return format(v, "f")
    return str(v)


class _Pdf(FPDF):
    def __init__(self, doc: Doc, rtl: bool, superseded_by: int | None) -> None:
        super().__init__(orientation="P", unit="mm", format="A4")
        self.doc = doc
        self.rtl = rtl
        self.superseded_by = superseded_by
        self.core_fonts_encoding = "cp1252"  # em dash and middle dot in the Latin footer
        self.alias_nb_pages()
        self.set_auto_page_break(auto=True, margin=16)
        self.add_font("dv", fname=str(FONTS / "DejaVuSans.ttf"))
        self.add_font("dv", style="B", fname=str(FONTS / "DejaVuSans-Bold.ttf"))
        _shaping(self, rtl)

    def header(self) -> None:
        _shaping(self, False)
        if self.superseded_by is not None:
            self.set_font("helvetica", "B", 14)
            self.set_text_color(200, 0, 0)
            self.cell(0, 8, f"SUPERSEDED BY Rev {self.superseded_by}", align="C",
                      new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
            self.set_text_color(0, 0, 0)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("helvetica", "", 7)
        d = self.doc
        self.cell(
            0, 6,
            f"{d.doc_no} · Rev {d.revision} · page {self.page_no()} of {{nb}} · "
            f"{d.snapshot_hash[:12]} · {d.footer_note}",
            align="C",
        )  # fmt: skip
        _shaping(self, self.rtl)


def ltr(v: str, rtl: bool) -> str:
    """Rule 42: numbers, codes and refs stay left-to-right inside RTL (LRI … PDI isolate)."""
    if rtl and any(c.isascii() and c.isalnum() for c in v):
        return f"\u2066{v}\u2069"
    return v


def _shaping(pdf: FPDF, rtl: bool) -> None:
    pdf.set_text_shaping(
        use_shaping_engine=True, direction="rtl" if rtl else "ltr",
        script="arab" if rtl else None, language="ara" if rtl else None,
    )  # fmt: skip


def _font(pdf: _Pdf, size: float, bold: bool = False) -> None:
    pdf.set_font("dv", "B" if bold else "", size)


def _table(pdf: _Pdf, t: Table, rtl: bool) -> None:
    cols = list(reversed(t.columns)) if rtl else t.columns
    if not cols:
        return
    title = t.title_ar if rtl else t.title_en
    if title:
        _font(pdf, 9, True)
        pdf.multi_cell(0, 5, title, align="R" if rtl else "L", new_x=XPos.LMARGIN,
                       new_y=YPos.NEXT)  # fmt: skip
    width = pdf.epw / len(cols)
    size = 7 if len(cols) <= 7 else 6
    _font(pdf, size, True)
    h = 5
    for _key, en, ar in cols:
        pdf.cell(width, h, (ar if rtl else en)[:40], border=1, align="C")
    pdf.ln(h)
    _font(pdf, size)
    for r in t.rows:
        if pdf.will_page_break(h):
            pdf.add_page()
        for key, _en, _ar in cols:
            pdf.cell(width, h, ltr(_txt(r.get(key))[:48], rtl), border=1, align="R" if rtl else "L")
        pdf.ln(h)
    pdf.ln(2)


def pdf(doc: Doc, lang: str, superseded_by: int | None = None) -> bytes:
    """`lang` = en, ar or both (EN then AR in one file)."""
    langs = ["en", "ar"] if lang == "both" else [lang]
    out: _Pdf | None = None
    for lg in langs:
        rtl = lg == "ar"
        if out is None:
            out = _Pdf(doc, rtl, superseded_by)
        else:
            out.rtl = rtl
            _shaping(out, rtl)
        p = out
        p.add_page()
        align = "R" if rtl else "L"
        _font(p, 15, True)
        p.multi_cell(0, 8, doc.title_ar if rtl else doc.title_en, align=align,
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
        _font(p, 8)
        for en, ar, v in doc.control:
            p.multi_cell(0, 4.5, f"{ar}: {ltr(v, True)}" if rtl else f"{en}: {v}", align=align,
                         new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
        for sec in doc.sections:
            _shaping(p, rtl)
            p.ln(2)
            _font(p, 11, True)
            p.multi_cell(0, 6, sec.title_ar if rtl else sec.title_en, align=align,
                         new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
            if sec.watermark:
                _shaping(p, False)
                p.set_font("helvetica", "B", 9)
                p.set_text_color(200, 0, 0)
                p.multi_cell(0, 5, sec.watermark.split(" / ")[0], align="C",
                             new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
                _shaping(p, rtl)
                _font(p, 9, True)
                p.multi_cell(0, 5, sec.watermark.split(" / ")[-1], align="C",
                             new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # fmt: skip
                p.set_text_color(0, 0, 0)
            _font(p, 8)
            for para in sec.paragraphs_ar if rtl else sec.paragraphs_en:
                p.multi_cell(0, 4.5, para, align=align, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            for t in sec.tables:
                _table(p, t, rtl)
    assert out is not None  # noqa: S101
    return bytes(out.output())


def _num(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        try:
            return float(Decimal(v.replace(",", "").replace(" %", "")))
        except Exception:
            return v
    return v


def xlsx(doc: Doc) -> bytes:
    """One sheet per table (EN and AR header rows; numbers as numeric cells; "—" as an empty
    cell with a note) and a "Document" sheet with the control data."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Document"
    ws0.append(["Field", "البيان", "Value"])
    ws0.append(["Document no.", "رقم الوثيقة", doc.doc_no])
    ws0.append(["Revision", "المراجعة", doc.revision])
    ws0.append(["Snapshot SHA-256", "بصمة البيانات", doc.snapshot_hash])
    for en, ar, v in doc.control:
        ws0.append([en, ar, v])
    used: set[str] = {"Document"}
    n = 0
    for sec in doc.sections:
        for t in sec.tables:
            n += 1
            base = (t.sheet or t.title_en or sec.title_en or f"Table {n}")[:28]
            name = base
            i = 2
            while name in used:
                name = f"{base[:26]}~{i}"
                i += 1
            used.add(name)
            ws = wb.create_sheet(name.replace("/", "-").replace(":", " "))
            ws.append([c[1] for c in t.columns])
            ws.append([c[2] for c in t.columns])
            for r in t.rows:
                row: list[Any] = []
                for key, _en, _ar in t.columns:
                    val = r.get(key)
                    row.append(None if val in (None, "", "—") else _num(val))
                ws.append(row)
                for idx, (key, _en, _ar) in enumerate(t.columns, start=1):
                    if r.get(key) in (None, "", "—"):
                        ws.cell(row=ws.max_row, column=idx).comment = Comment(
                            "— (no value / لا قيمة)", "HSE"
                        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
