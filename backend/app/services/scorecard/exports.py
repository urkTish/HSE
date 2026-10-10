"""The generic register export (spec 6g §3.8, EX-1…EX-11, SC-4): column selection by class,
rights and masks, purposes, CSV / XLSX / PDF encoding with the formula guard and the PDPL footer,
synchronous exports up to 5,000 rows and queued jobs above, expiring files in the encrypted bucket,
the export log, subscriptions and the dashboard PDF print."""

from __future__ import annotations

import csv
import hashlib
import io
import uuid
from datetime import datetime, timedelta
from typing import Any

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.clock import now
from app.core.enums import AuditAction, EntityType, ExportDataset, ExportFormat, NotificationKind
from app.core.errors import ApiError, ErrorCode, not_found, validation_error
from app.core.hse_enums import AttachmentOwner, ExportPurpose
from app.core.scorecard_enums import (
    XpFormat,
    XpFrequency,
    XpJobStatus,
    XpMaskMode,
    XpPdpl,
    XpPurpose,
    XpPurposeRule,
)
from app.models import Attachment, Incident, InjuryCase, User, XpJob, XpSubscription
from app.schemas.scorecard import (
    XpColumnRead,
    XpDatasetList,
    XpDatasetRead,
    XpExportRequest,
    XpFileUrl,
    XpJobPage,
    XpJobRead,
    XpSubscriptionCreate,
    XpSubscriptionList,
    XpSubscriptionRead,
)
from app.services import audit
from app.services.permissions import Grant, Principal, build_principal, deny, forbidden_error
from app.services.scorecard import common as cm
from app.services.scorecard import datasets as reg
from app.services.scorecard import render

SYNC_MAX = 5_000
JOB_MAX = 100_000
PDF_MAX = 2_000
DAILY_LIMIT = 50
CTYPES = {
    XpFormat.csv: "text/csv; charset=utf-8",
    XpFormat.xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    XpFormat.pdf: "application/pdf",
}
GUARD = ("=", "+", "-", "@")


class Out:
    """The produced table before encoding."""

    def __init__(self) -> None:
        self.cols: list[tuple[str, str, str]] = []
        self.rows: list[list[Any]] = []
        self.notes: list[str] = []
        self.personal = False
        self.sensitive: list[str] = []


# ---- registry reads ------------------------------------------------------------------------------


def _holds(p: Principal, project_id: uuid.UUID | None, cap: Any) -> bool:
    if cap is None or p.is_manager:
        return True
    if project_id is not None:
        return p.grant(project_id, cap) is not None
    return any(p.grant(pid, cap) is not None for pid in p.projects)


def dataset_read(p: Principal, ds: reg.Dataset) -> XpDatasetRead:
    return XpDatasetRead(
        dataset_code=ds.code, owner_module=ds.module, view_capability=ds.view_cap.value,
        export_capability=ds.export_cap.value, label_en=ds.en, label_ar=ds.ar,
        columns=[XpColumnRead(column_code=c.code, label_en=c.en, label_ar=c.ar, pdpl_class=c.pdpl,
                              needs_capability=c.needs.value if c.needs else None,
                              mask_mode=c.mask, default=c.default and c.pdpl == XpPdpl.none,
                              allowed=c.pdpl != XpPdpl.never and _holds(p, None, c.needs))
                 for c in reg.columns(ds)],
        purpose_required_when=ds.purpose_rule, aggregate=ds.aggregate,
        pdf_allowed=ds.pdf_allowed, project_required=ds.project_required,
    )  # fmt: skip


def datasets(db: Session, p: Principal) -> XpDatasetList:
    out = []
    for code in reg.all_codes():
        ds = reg.get(code)
        if ds is not None:
            out.append(dataset_read(p, ds))
    return XpDatasetList(items=out)


def _dataset(code: str) -> reg.Dataset:
    ds = reg.get(code)
    if ds is None:
        raise validation_error("dataset", "Unknown dataset.")
    return ds


# ---- column rules (EX-3, EX-4, EX-5) ------------------------------------------------------------


def _purpose_ok(purpose: XpPurpose | None, text: str | None) -> bool:
    return purpose is not None and (purpose != XpPurpose.other or len((text or "").strip()) >= 20)


def check_columns(
    p: Principal, ds: reg.Dataset, project_id: uuid.UUID | None, requested: list[str],
    purpose: XpPurpose | None, purpose_text: str | None,
) -> tuple[list[reg.Col], bool]:  # fmt: skip
    """Explicit columns: never → 422, no right → 403, sensitive without purpose → 422.
    Returns (declared columns requested, purpose needed)."""
    decl = {c.code: c for c in reg.columns(ds)}
    picked: list[reg.Col] = []
    need_purpose = ds.purpose_rule == XpPurposeRule.always
    for code in requested:
        c = decl.get(code)
        if c is None:
            continue
        if c.pdpl == XpPdpl.never:
            raise cm.code_err(ErrorCode.COLUMN_NOT_EXPORTABLE, f"{code} is never exported (EX-3).",
                              "هذا العمود لا يُصدَّر أبداً.", "columns", column=code)  # fmt: skip
        if c.pdpl in (XpPdpl.personal, XpPdpl.sensitive) and not _holds(p, project_id, c.needs):
            raise ApiError(403, ErrorCode.COLUMN_NOT_PERMITTED,
                           f"{code} needs {c.needs.value if c.needs else '—'} (EX-3).",
                           "لا تملك صلاحية تصدير هذا العمود.", meta={"column": code})  # fmt: skip
        if c.pdpl == XpPdpl.sensitive:
            need_purpose = True
        picked.append(c)
    if need_purpose and not _purpose_ok(purpose, purpose_text):
        raise cm.code_err(ErrorCode.PURPOSE_REQUIRED, "State the purpose of this export (EX-5).",
                          "حدد الغرض من التصدير.", "purpose")  # fmt: skip
    return picked, need_purpose


# ---- legacy registers ---------------------------------------------------------------------------


def legacy_csv(
    db: Session, p: Principal, code: str, project_id: uuid.UUID | None, status: str | None,
    identity: bool, purpose: XpPurpose | None, purpose_text: str | None,
) -> tuple[list[str], list[list[str]]]:  # fmt: skip
    ds = ExportDataset(code)
    fmt = ExportFormat.csv
    ep: ExportPurpose | None = None
    ptext = purpose_text
    if purpose is not None:
        try:
            ep = ExportPurpose(purpose.value)
        except ValueError:
            ep, ptext = ExportPurpose.other, f"{purpose.value}: {purpose_text or ''}".strip()
    with audit.muted(db, AuditAction.export):  # the job writes the one export entry
        content = _run_legacy(db, p, ds, fmt, project_id, status, identity, ep, ptext)
    text = content.decode("utf-8").lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(text)))
    return (rows[0] if rows else []), rows[1:]


def _run_legacy(
    db: Session, p: Principal, ds: ExportDataset, fmt: ExportFormat, project_id: uuid.UUID | None,
    status: str | None, identity: bool, ep: ExportPurpose | None, ptext: str | None,
) -> bytes:  # fmt: skip
    from app.services import exports as svc  # noqa: PLC0415
    from app.services import hse_exports  # noqa: PLC0415
    from app.services.access import exports as acc  # noqa: PLC0415
    from app.services.cert import exports as cer  # noqa: PLC0415
    from app.services.ptw import exports as ptw  # noqa: PLC0415
    from app.services.train import exports as trn  # noqa: PLC0415

    if ds in trn.DATASETS:
        content, _m, _f = trn.export(db, p, ds, fmt, project_id, status, None)
    elif ds in cer.DATASETS:
        content, _m, _f = cer.export(db, p, ds, fmt, project_id, status, None)
    elif ds in acc.DATASETS:
        content, _m, _f = acc.export(db, p, ds, fmt, project_id, status, None, identity, ep, ptext)
    elif ds.value in ("workforce_returns", "incidents", "observations", "inspections",
                      "corrective_actions", "hse_meetings"):  # fmt: skip
        content, _m, _f = hse_exports.export(db, p, ds, fmt, project_id, status, None, identity,
                                             ep, ptext)  # fmt: skip
    elif ds in ptw.DATASETS:
        content, _m, _f = ptw.export(db, p, ds, fmt, project_id, status, None)
    else:
        content, _m, _f = svc.export(db, p, ds, fmt, project_id, status, None)
    return content


def _privacy_cases(db: Session, project_id: uuid.UUID | None) -> set[str]:
    out: set[str] = set()
    for ref, no in db.execute(select(Incident.ref, InjuryCase.person_no)
                              .join(Incident, Incident.id == InjuryCase.incident_id)
                              .where(InjuryCase.privacy_case.is_(True),
                                     Incident.project_id == project_id)):  # fmt: skip
        out.add(f"{ref}-P{no}")
    return out


def _legacy(
    db: Session, p: Principal, ds: reg.Dataset, project_id: uuid.UUID | None,
    filters: dict[str, Any], requested: list[str], picked: list[reg.Col],
    purpose: XpPurpose | None, purpose_text: str | None,
) -> Out:  # fmt: skip
    sens = [c for c in picked if c.pdpl == XpPdpl.sensitive]
    header, rows = legacy_csv(db, p, ds.code, project_id, filters.get("status"), bool(sens),
                              purpose, purpose_text)  # fmt: skip
    decl = {c.code: c for c in ds.special}
    o = Out()
    derived = ["person"] if ds.code == "incidents" and "case_no" in header else []
    avail = header + derived
    if requested:
        unknown = [c for c in requested if c not in avail]
        if unknown:
            raise validation_error("columns", f"Unknown columns: {', '.join(unknown)}.")
        sel = list(requested)
    else:
        sel = [c for c in avail
               if c not in decl or decl[c].pdpl == XpPdpl.none]  # fmt: skip
        dropped = [c for c in header if c in decl and decl[c].pdpl != XpPdpl.none]
        if dropped:
            o.notes.append(f"Not in the de-identified default: {', '.join(dropped)}")
    idx = {c: i for i, c in enumerate(header)}
    g = p.grant(project_id, ds.export_cap) if project_id else None
    if g is not None and g.site_ids is not None and "site" in idx:  # site-scoped grant (EX-3)
        from app.models import Site  # noqa: PLC0415

        codes = set(db.scalars(select(Site.code).where(Site.id.in_(g.site_ids))))
        rows = [r for r in rows if r[idx["site"]] in codes]
    privacy = _privacy_cases(db, project_id) if ds.code == "incidents" else set()
    persons: dict[str, int] = {}
    for r in rows:
        out_row = []
        case = r[idx["case_no"]] if "case_no" in idx else ""
        for c in sel:
            if c == "person":
                if not case:
                    out_row.append("")
                    continue
                n = persons.setdefault(case, len(persons) + 1)
                trade = r[idx["trade"]] if "trade" in idx else ""
                emp = r[idx["employer"]] if "employer" in idx else ""
                out_row.append("Privacy case" if case in privacy else
                               f"Person {n} · {trade} · {emp}")  # fmt: skip
            elif c in ("person_name", "id_number", "id_type", "employee_no", "gosi_case_ref") \
                    and case in privacy:  # fmt: skip
                out_row.append("Privacy case" if c == "person_name" else "")
            else:
                out_row.append(r[idx[c]])
        o.rows.append(out_row)
    o.cols = [(c, decl[c].en if c in decl else reg.label_en(c),
               decl[c].ar if c in decl else reg.label_ar(c)) for c in sel]  # fmt: skip
    o.personal = any(c in decl and decl[c].pdpl == XpPdpl.personal for c in sel)
    o.sensitive = [c for c in sel if c in decl and decl[c].pdpl == XpPdpl.sensitive]
    return o


# ---- native registers ----------------------------------------------------------------------------


def _mask(
    c: reg.Col,
    val: Any,
    raw: Any,
    r: reg.Resolver,
    persons: dict[Any, int],
    db: Session,
    project_id: uuid.UUID | None,
) -> Any:
    if c.mask == XpMaskMode.person_n:
        if raw is None and not val:
            return ""
        n = persons.setdefault(raw or val, len(persons) + 1)
        return f"Person {n}"
    if c.mask == XpMaskMode.role_only:
        if raw is None or project_id is None:
            return "—" if val else ""
        roles = sorted(x.value for x in cm.roles_on(db, raw, project_id))
        return roles[0] if roles else "user"
    if c.mask == XpMaskMode.mask_id:
        s = str(val or "")
        return s[:1] + "*" * max(0, len(s) - 3) + s[-2:] if len(s) > 3 else s
    return ""


def _native(
    db: Session,
    p: Principal,
    ds: reg.Dataset,
    g: Grant | None,
    project_id: uuid.UUID | None,
    filters: dict[str, Any],
    requested: list[str],
    picked: list[reg.Col],
) -> Out:
    cols = reg.native_columns(ds)
    by = {c.code: c for c in cols}
    o = Out()
    if requested:
        unknown = [c for c in requested if c not in by]
        if unknown:
            raise validation_error("columns", f"Unknown columns: {', '.join(unknown)}.")
        sel = [by[c] for c in requested]
    else:
        sel = [c for c in cols if c.pdpl == XpPdpl.none or (
            c.pdpl == XpPdpl.personal and c.mask not in (None, XpMaskMode.omit))]  # fmt: skip
        dropped = [c.code for c in cols if c not in sel and c.pdpl != XpPdpl.never]
        if dropped:
            o.notes.append(f"Not in the de-identified default: {', '.join(dropped)}")
    objs = reg.native_rows(db, ds, g, project_id, filters)
    r = reg.Resolver(db)
    if ds.aggregate and project_id is not None and cm.is_viewer(p, project_id):
        return _aggregate(db, p, ds, project_id, objs, r)
    explicit = set(requested)
    persons: dict[Any, int] = {}
    for obj in objs:
        row = []
        for c in sel:
            val, raw = reg.native_value(r, obj, c)
            if c.pdpl == XpPdpl.personal and c.code not in explicit:
                val = _mask(c, val, raw, r, persons, db, project_id)
            row.append(val)
        o.rows.append(row)
    o.cols = [(c.code, c.en, c.ar) for c in sel]
    o.personal = any(c.pdpl == XpPdpl.personal and c.code in explicit for c in sel)
    o.sensitive = [c.code for c in sel if c.pdpl == XpPdpl.sensitive]
    return o


def _aggregate(db: Session, p: Principal, ds: reg.Dataset, project_id: uuid.UUID,
               objs: list[Any], r: reg.Resolver) -> Out:  # fmt: skip
    """EX-9: Viewer / Client get counts by month and contractor; 1–2 persons print "<3"."""
    small = not _holds(p, project_id, cm.C.breakdown_sensitive_view)
    ea = reg._eng_attr(ds.model)
    pcol = next((k for k in ("headcount", "persons_present") if hasattr(ds.model, k)), None)
    per_person = hasattr(ds.model, "worker_id")
    groups: dict[tuple[str, str], list[Any]] = {}
    for obj in objs:
        key = (reg.month_key(obj, ds), r.eng(getattr(obj, ea)) if ea else "")
        groups.setdefault(key, []).append(obj)
    o = Out()
    o.cols = [("month", "Month", "الشهر"), ("contractor", "Contractor", "المقاول"),
              ("records", "Records", "السجلات")]  # fmt: skip
    if pcol:
        o.cols.append(("persons", "Persons", "الأشخاص"))
    for (m, e), xs in sorted(groups.items()):
        n = len(xs)
        row: list[Any] = [m, e, "<3" if small and per_person and 0 < n < 3 else n]
        if pcol:
            s = sum(int(getattr(x, pcol) or 0) for x in xs)
            row.append("<3" if small and 0 < s < 3 else s)
        o.rows.append(row)
    o.notes.append("Aggregate rows only (Viewer / Client, EX-9).")
    return o


# ---- encoding (EX-6) ----------------------------------------------------------------------------


def guard(v: Any) -> Any:
    if isinstance(v, str) and v.startswith(GUARD):
        return "'" + v
    return v


def footer(p: Principal | User, export_no: str) -> str:
    u = p.user if isinstance(p, Principal) else p
    return (f"Exported {cm.to_local(now()):%Y-%m-%d %H:%M} by {u.full_name_en} · {export_no} · "
            "CONFIDENTIAL — PDPL")  # fmt: skip


def encode(o: Out, fmt: XpFormat, foot: str, title: tuple[str, str], export_no: str,
           lang: str) -> bytes:  # fmt: skip
    if fmt == XpFormat.csv:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([c[0] for c in o.cols])
        for r in o.rows:
            w.writerow([guard(reg.cell(v)) for v in r])
        return ("﻿" + buf.getvalue()).encode("utf-8")
    if fmt == XpFormat.xlsx:
        wb = Workbook()
        ws = wb.active
        ws.title = title[0][:31] or "Export"
        ws.append([c[1] for c in o.cols])
        ws.append([c[2] for c in o.cols])
        for r in o.rows:
            ws.append([guard(reg.cell(v)) for v in r])
        ws.append([])
        ws.append([foot])
        ws.oddFooter.center.text = foot
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    if len(o.rows) > PDF_MAX:
        raise validation_error("format", f"PDF exports hold at most {PDF_MAX:,} rows (EX-6).")
    doc = render.Doc(doc_no=export_no, revision=0, title_en=title[0], title_ar=title[1],
                     snapshot_hash=hashlib.sha256(repr(o.rows).encode()).hexdigest(),
                     control=[("Export", "التصدير", export_no)],
                     sections=[render.Section("data", title[0], title[1], [render.Table(
                         columns=o.cols, rows=[dict(zip([c[0] for c in o.cols], r, strict=False))
                                               for r in o.rows])])],
                     footer_note=foot.split(" · ", 1)[-1])  # fmt: skip
    return render.pdf(doc, lang)


# ---- jobs ----------------------------------------------------------------------------------------


def _export_no(db: Session, t: datetime) -> tuple[int, int, str]:
    year = cm.to_local(t).year
    seq = (db.scalar(select(func.max(XpJob.seq)).where(XpJob.year == year)) or 0) + 1
    return year, seq, f"EXP-{year}-{seq:06d}"


def _scope(
    db: Session, p: Principal, ds: reg.Dataset, project_id: uuid.UUID | None
) -> Grant | None:
    if project_id is None:
        if ds.project_required:
            raise validation_error("project_id", "Choose a project to export.")
        if not ds.legacy and not _holds(p, None, ds.export_cap):
            raise forbidden_error()
        return None
    cm.project(db, p, project_id)
    if ds.legacy:
        return None
    g = p.grant(project_id, ds.view_cap)
    if g is None or (not p.is_manager and p.grant(project_id, ds.export_cap) is None
                     and not cm.is_viewer(p, project_id)):  # fmt: skip
        raise forbidden_error(f"Exporting {ds.code} needs {ds.export_cap.value}.")
    return g


def produce(
    db: Session, p: Principal, ds: reg.Dataset, project_id: uuid.UUID | None,
    filters: dict[str, Any], requested: list[str], purpose: XpPurpose | None,
    purpose_text: str | None,
) -> Out:  # fmt: skip
    g = _scope(db, p, ds, project_id)
    picked, _need = check_columns(p, ds, project_id, requested, purpose, purpose_text)
    if ds.legacy:
        return _legacy(db, p, ds, project_id, filters, requested, picked, purpose, purpose_text)
    return _native(db, p, ds, g, project_id, filters, requested, picked)


def _rate_limit(db: Session, p: Principal, t: datetime) -> None:
    start = cm.at_local(cm.local_day(t), "00:00")
    n = db.scalar(select(func.count()).select_from(XpJob).where(
        XpJob.requested_by_user_id == p.user.id, XpJob.created_at >= start,
        XpJob.subscription_id.is_(None)))  # fmt: skip
    if (n or 0) >= DAILY_LIMIT:
        raise ApiError(429, ErrorCode.EXPORT_RATE_LIMIT,
                       f"At most {DAILY_LIMIT} exports per day (EX-7).",
                       "الحد الأقصى 50 عملية تصدير يومياً.")  # fmt: skip


def _store(db: Session, p: Principal | User, job: XpJob, ds: reg.Dataset, o: Out,
           t: datetime) -> bytes:  # fmt: skip
    from app.services import attachments  # noqa: PLC0415

    u = p.user if isinstance(p, Principal) else p
    lang = u.preferred_language.value if u.preferred_language else "en"
    content = encode(o, job.format, footer(p, job.export_no), (ds.en, ds.ar), job.export_no, lang)
    pid = job.project_id or db.scalar(select(Attachment.project_id).limit(1))
    if pid is None:
        from app.models import Project  # noqa: PLC0415

        pid = db.scalar(select(Project.id).limit(1))
    assert pid is not None  # noqa: S101
    ext = job.format.value
    a = attachments.store(
        db,
        AttachmentOwner.xp_export_file,
        job.id,
        pid,
        f"{ds.code}-{job.export_no}.{ext}",
        content,
        CTYPES[job.format],
        u.id,
    )
    job.attachment_id = a.id
    job.sha256 = a.sha256
    job.row_count = len(o.rows)
    job.columns = [c[0] for c in o.cols]
    job.notes = list(o.notes)
    job.contains_personal = o.personal or bool(o.sensitive)
    job.contains_sensitive = bool(o.sensitive)
    job.status = XpJobStatus.ready
    job.ready_at = t
    job.expires_at = t + (timedelta(hours=24) if o.sensitive else timedelta(days=7))
    actor = p.actor(job.project_id) if isinstance(p, Principal) else audit.SYSTEM
    audit.record(db, AuditAction.export, actor, entity_type=EntityType.export_job,
                 entity_id=job.id, project_id=job.project_id,
                 details={"dataset": job.dataset, "filters": job.filters, "columns": job.columns,
                          "row_count": job.row_count,
                          "purpose": job.purpose.value if job.purpose else None,
                          "purpose_text": job.purpose_text, "sha256": job.sha256,
                          "export_id": job.export_no, "format": job.format.value},
                 fields_read=o.sensitive or None)  # fmt: skip
    if o.sensitive:
        audit.record(db, AuditAction.sensitive_field_read, actor,
                     entity_type=EntityType.export_job, entity_id=job.id,
                     project_id=job.project_id, fields_read=o.sensitive,
                     details={"export_id": job.export_no,
                              "purpose": job.purpose.value if job.purpose else None})  # fmt: skip
    return content


def new_job(
    db: Session, p: Principal, body: XpExportRequest | XpSubscription, t: datetime,
    subscription_id: uuid.UUID | None = None,
) -> XpJob:  # fmt: skip
    return _new_job(db, p, body, t, subscription_id)[0]


def _new_job(
    db: Session, p: Principal, body: XpExportRequest | XpSubscription, t: datetime,
    subscription_id: uuid.UUID | None = None,
) -> tuple[XpJob, bytes | None]:  # fmt: skip
    ds = _dataset(body.dataset)
    fmt = XpFormat(body.format)
    if fmt == XpFormat.pdf and not ds.pdf_allowed:
        raise validation_error("format", "PDF is not offered for this dataset.")
    purpose = getattr(body, "purpose", None)
    ptext = getattr(body, "purpose_text", None)
    o = produce(db, p, ds, body.project_id, dict(body.filters or {}), list(body.columns or []),
                purpose, ptext)  # fmt: skip
    if len(o.rows) > JOB_MAX:
        raise cm.code_err(ErrorCode.EXPORT_TOO_LARGE,
                          f"More than {JOB_MAX:,} rows: narrow the filters (EX-7).",
                          "عدد الصفوف كبير جداً؛ ضيّق عوامل التصفية.")  # fmt: skip
    year, seq, no = _export_no(db, t)
    job = XpJob(id=uuid.uuid4(), year=year, seq=seq, export_no=no, dataset=ds.code, format=fmt,
                project_id=body.project_id, filters=dict(body.filters or {}),
                columns=list(body.columns or []), notes=list(o.notes), purpose=purpose,
                purpose_text=ptext, status=XpJobStatus.queued, subscription_id=subscription_id,
                requested_by_user_id=p.user.id, created_at=t)  # fmt: skip
    db.add(job)
    db.flush()
    content = _store(db, p, job, ds, o, t) if len(o.rows) <= SYNC_MAX else None
    db.flush()
    return job, content


def request(db: Session, p: Principal, body: XpExportRequest) -> XpJobRead:
    t = now()
    _rate_limit(db, p, t)
    return to_read(db, new_job(db, p, body, t))


def run_queue(db: Session, t: datetime) -> int:
    """`export_jobs`: queued jobs are produced again as the requester and stored (EX-7)."""
    n = 0
    for job in db.scalars(select(XpJob).where(XpJob.status == XpJobStatus.queued)):
        u = db.get(User, job.requested_by_user_id)
        if u is None:
            continue
        p = build_principal(db, u, None)
        ds = _dataset(job.dataset)
        try:
            o = produce(db, p, ds, job.project_id, dict(job.filters or {}), list(job.columns or []),
                        job.purpose, job.purpose_text)  # fmt: skip
            _store(db, p, job, ds, o, t)
            cm.send(
                db,
                [u.id],
                NotificationKind.export_ready,
                f"Export {job.export_no} ({job.dataset}) is ready",
                f"التصدير {job.export_no} جاهز",
                job.project_id,
                EntityType.export_job,
                job.id,
            )
        except ApiError as e:
            job.status, job.error = XpJobStatus.failed, e.message[:300]
            cm.send(db, [u.id], NotificationKind.export_ready,
                    f"Export {job.export_no} failed: {job.error}", f"فشل التصدير {job.export_no}",
                    job.project_id, EntityType.export_job, job.id)  # fmt: skip
        n += 1
    db.flush()
    return n


def purge(db: Session, t: datetime) -> int:
    """`export_purge` 03:00: files past expires_at are deleted and the job becomes expired."""
    from app.services import attachments  # noqa: PLC0415

    n = 0
    for job in db.scalars(select(XpJob).where(XpJob.status == XpJobStatus.ready,
                                              XpJob.expires_at <= t)):  # fmt: skip
        a = db.get(Attachment, job.attachment_id) if job.attachment_id else None
        if a is not None:
            job.attachment_id = None
            db.flush()
            attachments.erase(db, a)
        job.status = XpJobStatus.expired
        n += 1
    db.flush()
    return n


# ---- reads ---------------------------------------------------------------------------------------


def to_read(db: Session, j: XpJob) -> XpJobRead:
    return XpJobRead(
        id=j.id, export_no=j.export_no, dataset=j.dataset, format=j.format,
        project_id=j.project_id, filters=j.filters or {}, columns=list(j.columns or []),
        notes=list(j.notes or []), purpose=j.purpose, purpose_text=j.purpose_text,
        row_count=j.row_count, sha256=j.sha256, contains_personal=j.contains_personal,
        contains_sensitive=j.contains_sensitive, status=j.status, expires_at=j.expires_at,
        requested_by=cm.user_ref(db, j.requested_by_user_id), created_at=j.created_at,
        ready_at=j.ready_at,
    )  # fmt: skip


def _visible(db: Session, p: Principal, j: XpJob) -> bool:
    """EX-10: own jobs; HSE Officers their projects' jobs; the HSE Manager all."""
    if p.is_manager or j.requested_by_user_id == p.user.id:
        return True
    return j.project_id is not None and cm.staff(db, p, j.project_id)


def list_jobs(
    db: Session, p: Principal, project_id: uuid.UUID | None, dataset: str | None,
    status: XpJobStatus | None, mine: bool, page: int, size: int,
) -> XpJobPage:  # fmt: skip
    stmt = select(XpJob)
    if project_id is not None:
        stmt = stmt.where(XpJob.project_id == project_id)
    if dataset:
        stmt = stmt.where(XpJob.dataset == dataset)
    if status is not None:
        stmt = stmt.where(XpJob.status == status)
    if mine:
        stmt = stmt.where(XpJob.requested_by_user_id == p.user.id)
    rows = [j for j in db.scalars(stmt.order_by(XpJob.created_at.desc())) if _visible(db, p, j)]
    return XpJobPage(items=[to_read(db, j) for j in rows[(page - 1) * size : page * size]],
                     total=len(rows), page=page, page_size=size)  # fmt: skip


def read(db: Session, p: Principal, job_id: uuid.UUID) -> XpJobRead:
    j = db.get(XpJob, job_id)
    if j is None or not _visible(db, p, j):
        raise deny(db, p, EntityType.export_job, job_id, j.project_id if j else None, "Export")
    return to_read(db, j)


def file_url(db: Session, p: Principal, job_id: uuid.UUID) -> XpFileUrl:
    """EX-8: only the requester downloads, through a signed URL valid ≤ 5 min."""
    from app.services import attachments  # noqa: PLC0415

    j = db.get(XpJob, job_id)
    if j is None or j.requested_by_user_id != p.user.id:
        raise deny(db, p, EntityType.export_job, job_id, j.project_id if j else None, "Export")
    if j.status != XpJobStatus.ready or j.attachment_id is None or (
            j.expires_at is not None and j.expires_at <= now()):  # fmt: skip
        raise not_found("Export file")
    ttl = 300
    return XpFileUrl(url=attachments.raw_signed_url(j.attachment_id, ttl),
                     expires_at=now() + timedelta(seconds=ttl))  # fmt: skip


# ---- subscriptions (SC-4) ------------------------------------------------------------------------


def _sub_read(s: XpSubscription) -> XpSubscriptionRead:
    return XpSubscriptionRead(
        id=s.id,
        dataset=s.dataset,
        project_id=s.project_id,
        filters=s.filters or {},
        columns=list(s.columns or []),
        format=s.format,
        frequency=s.frequency,
        active=s.active,
        last_run_at=s.last_run_at,
    )


def list_subscriptions(db: Session, p: Principal) -> XpSubscriptionList:
    rows = db.scalars(
        select(XpSubscription).where(
            XpSubscription.user_id == p.user.id, XpSubscription.active.is_(True)
        )
    )
    return XpSubscriptionList(items=[_sub_read(s) for s in rows])


def subscribe(db: Session, p: Principal, body: XpSubscriptionCreate) -> XpSubscriptionRead:
    ds = _dataset(body.dataset)
    _scope(db, p, ds, body.project_id)
    decl = {c.code: c for c in reg.columns(ds)}
    if any(c in decl and decl[c].pdpl != XpPdpl.none for c in body.columns):
        raise cm.code_err(
            ErrorCode.SUBSCRIPTION_PERSONAL_DATA,
            "Subscriptions cannot hold personal or sensitive columns (§3.8).",
            "لا يمكن الاشتراك في أعمدة شخصية.",
            "columns",
        )
    s = XpSubscription(
        id=uuid.uuid4(),
        user_id=p.user.id,
        project_id=body.project_id,
        dataset=ds.code,
        filters=body.filters,
        columns=body.columns,
        format=body.format,
        frequency=body.frequency,
        active=True,
    )
    db.add(s)
    db.flush()
    cm.record(db, p, AuditAction.create, EntityType.export_subscription, s, body.project_id)
    return _sub_read(s)


def unsubscribe(db: Session, p: Principal, sub_id: uuid.UUID) -> None:
    s = db.get(XpSubscription, sub_id)
    if s is None or s.user_id != p.user.id:
        raise not_found("Subscription")
    s.active = False
    cm.record(db, p, AuditAction.update, EntityType.export_subscription, s, s.project_id)
    db.flush()


def run_subscriptions(db: Session, t: datetime) -> int:
    """`export_subscriptions` 05:30: weekly on Sunday, monthly on day 1 (SC-4)."""
    today = cm.local_day(t)
    n = 0
    for s in db.scalars(select(XpSubscription).where(XpSubscription.active.is_(True))):
        due = (s.frequency == XpFrequency.weekly and today.weekday() == 6) or (
            s.frequency == XpFrequency.monthly and today.day == 1)  # fmt: skip
        if not due or (s.last_run_at is not None and cm.local_day(s.last_run_at) == today):
            continue
        u = db.get(User, s.user_id)
        if u is None:
            continue
        p = build_principal(db, u, None)
        try:
            job = new_job(db, p, s, t, s.id)
        except ApiError:
            continue
        s.last_run_at = t
        cm.send(db, [u.id], NotificationKind.export_ready,
                f"Scheduled export {job.export_no} ({s.dataset}) is in your export list",
                f"التصدير المجدول {job.export_no} جاهز", s.project_id, EntityType.export_job,
                job.id)  # fmt: skip
        n += 1
    db.flush()
    return n


# ---- GET /exports/{dataset} for registry-native datasets and the dashboard print ---------------


def sync_get(
    db: Session, p: Principal, code: str, fmt: ExportFormat, project_id: uuid.UUID | None,
    status: str | None,
) -> tuple[bytes, str, str]:  # fmt: skip
    t = now()
    _rate_limit(db, p, t)
    body = XpExportRequest(dataset=code, format=XpFormat(fmt.value), project_id=project_id,
                           filters={"status": status} if status else {})  # fmt: skip
    job, content = _new_job(db, p, body, t)
    if content is None:
        raise cm.code_err(ErrorCode.EXPORT_TOO_LARGE,
                          f"More than {SYNC_MAX:,} rows: use POST /exports (EX-7).",
                          "استخدم التصدير غير المتزامن.")  # fmt: skip
    return content, CTYPES[job.format], f"{code}-{job.export_no}.{job.format.value}"


def dashboard_pdf(db: Session, p: Principal, q: Any) -> tuple[bytes, str]:
    """EX-11 / D-10: the KPI endpoint values through the RP-5 renderer; an `export` audit row."""
    from app.kpi import scope as kscope  # noqa: PLC0415
    from app.kpi import service  # noqa: PLC0415

    sc = kscope.build(db, p, q)
    vals = service.metric_list(sc, None)
    t = now()
    no = f"DASH-{cm.to_local(t):%Y%m%d%H%M}"
    rows = [
        {
            "kpi": v.metric.value,
            "label": v.label_en,
            "value": v.display,
            "numerator": v.numerator or "—",
            "denominator": v.denominator or "—",
        }
        for v in vals
    ]
    foot = footer(p, no)
    doc = render.Doc(doc_no=no, revision=0, title_en="HSE KPI dashboard",
                     title_ar="لوحة مؤشرات الصحة والسلامة",
                     snapshot_hash=hashlib.sha256(repr(rows).encode()).hexdigest(),
                     control=[("Scope", "النطاق", service.scope_label(sc))],
                     sections=[render.Section("kpis", "KPIs", "المؤشرات", [render.Table(
                         [("kpi", "KPI", "المؤشر"), ("label", "Name", "الاسم"),
                          ("value", "Value", "القيمة"), ("numerator", "Numerator", "البسط"),
                          ("denominator", "Denominator", "المقام")], rows)])],
                     footer_note=foot.split(" · ", 1)[-1])  # fmt: skip
    lang = p.user.preferred_language.value if p.user.preferred_language else "en"
    content = render.pdf(doc, lang)
    audit.record(
        db,
        AuditAction.export,
        p.actor(None),
        details={
            "dataset": "dashboard_pdf",
            "row_count": len(rows),
            "export_id": no,
            "sha256": hashlib.sha256(content).hexdigest(),
            "filters": {"scope": service.scope_label(sc)},
        },
    )
    return content, f"{no}.pdf"
