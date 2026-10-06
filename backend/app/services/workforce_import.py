"""Workforce CSV/Excel import with dry run (spec 1-dashboard §3.2, rules W-5…W-7, E01-E14,
W01-W06)."""

import csv
import hashlib
import io
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.clock import now
from app.core.enums import (
    AuditAction,
    Capability,
    ContractorStatus,
    EntityType,
    ExportFormat,
    NotificationKind,
    ZoneStatus,
)
from app.core.errors import ApiError, ErrorCode
from app.core.hse_enums import (
    ImportCode,
    ImportMode,
    ImportRowStatus,
    ImportStatus,
    Shift,
    WorkforceSource,
    WorkforceStatus,
)
from app.kpi.periods import month_start
from app.models import (
    PeriodLock,
    Project,
    ProjectEngagement,
    Site,
    WorkforceImportBatch,
    WorkforceReturn,
    Zone,
)
from app.schemas.hse_common import UserRef
from app.schemas.workforce import (
    ImportCounts,
    ImportIssue,
    ImportRowReport,
    WorkforceImportPage,
    WorkforceImportRead,
    WorkforceImportSummary,
)
from app.services import audit, hse_settings, notify, projects
from app.services import exports as export_svc
from app.services import workforce as wf
from app.services.common import ensure_open, paginate
from app.services.hse_common import Refs, project_today
from app.services.permissions import Principal, deny, forbidden_error

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
EXPIRY = timedelta(minutes=60)
Code = ImportCode

COLUMNS: list[tuple[str, str, bool]] = [  # (key, Arabic header, required)
    ("work_date", "تاريخ_العمل", True),
    ("project_code", "رمز_المشروع", True),
    ("site_code", "رمز_الموقع", True),
    ("zone_code", "رمز_المنطقة", False),
    ("contractor_code", "رمز_المقاول", True),
    ("shift", "الوردية", True),
    ("headcount", "عدد_العاملين", True),
    ("man_hours", "ساعات_العمل", True),
    ("no_work", "لا_يوجد_عمل", False),
    ("toolbox_talks", "اجتماعات_التوعية", False),
    ("toolbox_attendees", "حضور_التوعية", False),
    ("inductions", "التعريفات", False),
    ("training_hours", "ساعات_التدريب", False),
    ("remarks", "ملاحظات", False),
]
HEADER_MAP = {k: k for k, _, _ in COLUMNS} | {ar: k for k, ar, _ in COLUMNS}
SHIFT_MAP = {
    "day": Shift.day,
    "night": Shift.night,
    "all": Shift.all,
    "نهارية": Shift.day,
    "ليلية": Shift.night,
    "كامل": Shift.all,
    "كامل اليوم": Shift.all,
}
YES = {"y", "yes", "1", "true", "نعم"}
NO = {"", "n", "no", "0", "false", "لا"}
MESSAGES: dict[ImportCode, tuple[str, str]] = {
    Code.E01: (
        "Unknown project code or not this project",
        "رمز مشروع غير معروف أو ليس هذا المشروع",
    ),
    Code.E02: ("Site code not in project", "رمز الموقع غير موجود في المشروع"),
    Code.E03: ("Zone code not in site, or zone archived", "رمز المنطقة ليس ضمن الموقع أو مؤرشف"),
    Code.E04: (
        "Contractor not engaged on the project, or site not in its sites",
        "المقاول غير مرتبط بالمشروع أو الموقع ليس ضمن مواقعه",
    ),
    Code.E05: (
        "Date outside the engagement mobilisation–demobilisation",
        "التاريخ خارج فترة تعبئة المقاول",
    ),
    Code.E06: (
        "Date in the future or before the project start",
        "التاريخ في المستقبل أو قبل بداية المشروع",
    ),
    Code.E07: (
        "Headcount must be an integer 0–20,000",
        "عدد العاملين يجب أن يكون عدداً صحيحاً 0–20,000",
    ),
    Code.E08: (
        "Man-hours negative, not a number, or above headcount × maximum hours",
        "ساعات العمل سالبة أو غير رقمية أو أعلى من الحد الأقصى",
    ),
    Code.E09: ("Duplicate key within the file", "مفتاح مكرر داخل الملف"),
    Code.E10: ("Row already exists (insert only)", "السجل موجود مسبقاً (إضافة فقط)"),
    Code.E11: ("Month is locked", "الشهر مقفل"),
    Code.E12: ("Row outside your scope (site/contractor)", "السجل خارج نطاق صلاحياتك"),
    Code.E13: (
        "Headcount 0 with man-hours, or no work with values",
        "عدد العاملين صفر مع ساعات، أو لا يوجد عمل مع قيم",
    ),
    Code.E14: (
        "Unparseable date or number, or missing required column",
        "تاريخ أو رقم غير صالح، أو عمود مطلوب مفقود",
    ),
    Code.W01: (
        "More hours per person than the warning threshold",
        "ساعات للفرد أعلى من حد التنبيه",
    ),
    Code.W02: (
        "Headcount differs > 50 % from the 7-day average",
        "عدد العاملين يختلف بأكثر من 50 % عن متوسط 7 أيام",
    ),
    Code.W03: ("Contractor suspended (hours still accepted)", "المقاول موقوف (تُقبل الساعات)"),
    Code.W04: ("Work date older than 30 days", "تاريخ العمل أقدم من 30 يوماً"),
    Code.W05: ("The same file was already committed on this project", "تم اعتماد نفس الملف مسبقاً"),
    Code.W06: (
        "Both an 'all' shift and a day/night row exist for this key",
        "يوجد سجل وردية كاملة وسجل نهاري/ليلي لنفس المفتاح",
    ),
}


def issue(code: ImportCode, column: str | None = None) -> dict[str, Any]:
    en, ar = MESSAGES[code]
    return {"code": code.value, "column": column, "message_en": en, "message_ar": ar}


def file_invalid(message: str) -> ApiError:
    return ApiError(422, ErrorCode.IMPORT_FILE_INVALID, message, "ملف الاستيراد غير صالح.")


# ---- parsing -------------------------------------------------------------------------------------


def _read_table(name: str, content: bytes) -> list[list[Any]]:
    lower = name.lower()
    if lower.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise file_invalid("The Excel file could not be read.") from exc
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if lower.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise file_invalid("CSV files must be UTF-8.") from exc
        first = text.splitlines()[0] if text else ""
        delim = ";" if first.count(";") > first.count(",") else ","
        return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    raise ApiError(
        415,
        ErrorCode.FILE_TYPE_NOT_ALLOWED,
        "Upload a .csv or .xlsx file.",
        "يرجى رفع ملف ‎.csv أو ‎.xlsx.",
    )


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", "_", str(h or "").strip()).lower()


def _parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    for fmt_ in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt_).date()
        except ValueError:
            continue
    return None


def _parse_number(v: Any) -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return Decimal(0)
    if isinstance(v, int | float | Decimal) and not isinstance(v, bool):
        return Decimal(str(v))
    s = str(v).strip()
    if not re.fullmatch(r"-?\d+(\.\d+)?", s):  # thousands separators rejected
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def parse(name: str, content: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """→ (rows with raw/parsed values and row-level parse issues, file-level issues)."""
    if len(content) > MAX_BYTES:
        raise file_invalid("The file is larger than 5 MB.")
    table = _read_table(name, content)
    while table and all(c in (None, "") for c in table[-1]):
        table.pop()
    if not table:
        raise file_invalid("The file is empty.")
    headers = [
        HEADER_MAP.get(_norm_header(h)) or HEADER_MAP.get(str(h or "").strip()) for h in table[0]
    ]
    missing = [k for k, _, req in COLUMNS if req and k not in headers]
    if missing:
        raise ApiError(
            422,
            ErrorCode.IMPORT_FILE_INVALID,
            "Missing required column(s): " + ", ".join(missing) + " (E14).",
            "أعمدة مطلوبة مفقودة: " + "، ".join(missing),
        )
    data = [r for r in table[1:] if any(c not in (None, "") for c in r)]
    if len(data) > MAX_ROWS:
        raise file_invalid(f"The file has more than {MAX_ROWS:,} data rows.")
    rows = []
    for i, raw in enumerate(data, start=1):
        vals = {h: raw[j] if j < len(raw) else None for j, h in enumerate(headers) if h}
        row: dict[str, Any] = {"row_no": i, "issues": []}
        d = _parse_date(vals.get("work_date"))
        if d is None:
            row["issues"].append(issue(Code.E14, "work_date"))
        row["work_date"] = d.isoformat() if d else None
        for k in ("project_code", "site_code", "zone_code", "contractor_code", "remarks"):
            v = vals.get(k)
            row[k] = str(v).strip() if v not in (None, "") else None
        sh = str(vals.get("shift") or "").strip().lower()
        shift = SHIFT_MAP.get(sh) or SHIFT_MAP.get(str(vals.get("shift") or "").strip())
        if shift is None:
            row["issues"].append(issue(Code.E14, "shift"))
        row["shift"] = shift.value if shift else None
        nw = str(vals.get("no_work") or "").strip().lower()
        if nw in YES:
            row["no_work"] = True
        elif nw in NO:
            row["no_work"] = False
        else:
            row["no_work"] = False
            row["issues"].append(issue(Code.E14, "no_work"))
        hc = _parse_number(vals.get("headcount"))
        if hc is None:
            row["issues"].append(issue(Code.E14, "headcount"))
        elif hc != hc.to_integral_value() or not 0 <= hc <= 20_000:
            row["issues"].append(issue(Code.E07, "headcount"))
            hc = None
        row["headcount"] = int(hc) if hc is not None else None
        mh = _parse_number(vals.get("man_hours"))
        if mh is None:
            row["issues"].append(issue(Code.E08, "man_hours"))
        row["man_hours"] = str(mh) if mh is not None else None
        for k in ("toolbox_talks", "toolbox_attendees", "inductions"):
            n = _parse_number(vals.get(k))
            if n is None or n < 0 or n != n.to_integral_value():
                row["issues"].append(issue(Code.E14, k))
                n = Decimal(0)
            row[k] = int(n)
        th = _parse_number(vals.get("training_hours"))
        if th is None or th < 0:
            row["issues"].append(issue(Code.E14, "training_hours"))
            th = Decimal(0)
        row["training_hours"] = str(th)
        rows.append(row)
    return rows, []


# ---- validation ----------------------------------------------------------------------------------


def validate(
    db: Session,
    p: Principal,
    project: Project,
    rows: list[dict[str, Any]],
    mode: ImportMode,
    sha: str,
    exclude_batch: uuid.UUID | None = None,
) -> list[dict[str, Any]]:
    """Adds `codes`, `status`, `action`, resolved ids to each row (in place) and returns them."""
    s = hse_settings.get(db, project.id)
    today = project_today(project)
    grant = p.grant(project.id, Capability.workforce_import)
    sites = {
        x.code.lower(): x for x in db.scalars(select(Site).where(Site.project_id == project.id))
    }
    zones: dict[tuple[uuid.UUID, str], Zone] = {
        (z.site_id, z.code.lower()): z
        for z in db.scalars(select(Zone).where(Zone.project_id == project.id))
    }
    engs = {
        e.contractor.short_code.lower(): e
        for e in db.scalars(
            select(ProjectEngagement).where(ProjectEngagement.project_id == project.id)
        )
    }
    locked = {
        x.month
        for x in db.scalars(
            select(PeriodLock).where(
                PeriodLock.project_id == project.id, PeriodLock.locked.is_(True)
            )
        )
    }
    W = WorkforceReturn  # noqa: N806
    sha_committed = db.scalar(
        select(func.count())
        .select_from(WorkforceImportBatch)
        .where(
            WorkforceImportBatch.project_id == project.id,
            WorkforceImportBatch.file_sha256 == sha,
            WorkforceImportBatch.status == ImportStatus.committed,
            WorkforceImportBatch.id != (exclude_batch or uuid.uuid4()),
        )
    )
    seen: dict[tuple[Any, ...], int] = {}
    shifts_by_key: dict[tuple[Any, ...], set[str]] = defaultdict(set)
    # pre-load existing keys for the dates in the file
    dates = sorted({r["work_date"] for r in rows if r.get("work_date")})
    existing: dict[tuple[Any, ...], WorkforceReturn] = {}
    hc_hist: dict[tuple[uuid.UUID, uuid.UUID], list[tuple[date, int]]] = defaultdict(list)
    if dates:
        lo, hi = date.fromisoformat(dates[0]), date.fromisoformat(dates[-1])
        for wr in db.scalars(
            select(W).where(W.project_id == project.id, W.work_date >= lo, W.work_date <= hi)
        ):
            k5 = (wr.work_date.isoformat(), wr.site_id, wr.zone_id, wr.engagement_id)
            existing[(*k5, wr.shift.value)] = wr
            shifts_by_key[k5].add(wr.shift.value)
        for d_, eng_id, site_id, hc_sum in db.execute(
            select(W.work_date, W.engagement_id, W.site_id, func.sum(W.headcount))
            .where(
                W.project_id == project.id,
                W.work_date >= lo - timedelta(days=7),
                W.work_date < hi,
                W.status != WorkforceStatus.draft,
            )
            .group_by(W.work_date, W.engagement_id, W.site_id)
        ):
            hc_hist[(eng_id, site_id)].append((d_, int(hc_sum or 0)))
    file_shifts: dict[tuple[Any, ...], set[str]] = defaultdict(set)
    for r in rows:
        if r.get("work_date") and r.get("shift"):
            site = sites.get((r.get("site_code") or "").lower())
            eng = engs.get((r.get("contractor_code") or "").lower())
            zone = (
                zones.get((site.id, (r.get("zone_code") or "").lower()))
                if site and r.get("zone_code")
                else None
            )
            file_shifts[
                (
                    r["work_date"],
                    site.id if site else None,
                    zone.id if zone else None,
                    eng.id if eng else None,
                )
            ].add(r["shift"])
    for r in rows:
        issues: list[dict[str, Any]] = list(r.get("parse_issues", r["issues"]))
        r["parse_issues"] = list(issues)
        codes = {i["code"] for i in issues}

        def add(
            code: ImportCode,
            column: str | None = None,
            codes: set[str] = codes,
            issues: list[dict[str, Any]] = issues,
        ) -> None:
            if code.value not in codes:
                codes.add(code.value)
                issues.append(issue(code, column))

        if (r.get("project_code") or "").lower() != project.code.lower():
            add(Code.E01, "project_code")
        site = sites.get((r.get("site_code") or "").lower())
        if site is None:
            add(Code.E02, "site_code")
        zone = None
        if r.get("zone_code"):
            zone = zones.get((site.id, r["zone_code"].lower())) if site else None
            if zone is None or zone.status == ZoneStatus.archived:
                add(Code.E03, "zone_code")
        eng = engs.get((r.get("contractor_code") or "").lower())
        if eng is None or (site is not None and site.id not in (eng.site_ids or [])):
            add(Code.E04, "contractor_code")
        d = date.fromisoformat(r["work_date"]) if r.get("work_date") else None
        if (
            d
            and eng
            and (
                d < eng.mobilisation_date
                or (eng.demobilisation_date and d > eng.demobilisation_date)
            )
        ):
            add(Code.E05, "work_date")
        if d and (d > today or d < project.start_date):
            add(Code.E06, "work_date")
        hc = r.get("headcount")
        mh = Decimal(r["man_hours"]) if r.get("man_hours") is not None else None
        if mh is not None and (mh < 0 or (hc is not None and mh > hc * s.max_hours_per_person_day)):
            add(Code.E08, "man_hours")
        if (
            hc is not None
            and mh is not None
            and ((hc == 0 and mh > 0) or (r.get("no_work") and (hc > 0 or mh > 0)))
        ):
            add(Code.E13, "headcount")
        if (
            grant is None
            or (site and not grant.covers_site(site.id))
            or (eng and grant.engagement_ids is not None and eng.id not in grant.engagement_ids)
            or (eng is None and grant.engagement_ids is not None)
        ):
            add(Code.E12)
        key = (
            r.get("work_date"),
            site.id if site else None,
            zone.id if zone else None,
            eng.id if eng else None,
            r.get("shift"),
        )
        if all(k is not None for k in (key[0], key[1], key[3], key[4])):
            if key in seen:
                add(Code.E09)
            seen.setdefault(key, r["row_no"])
        r["action"] = "insert"
        ex = (
            existing.get(key)
            if all(x is not None for x in (key[0], key[1], key[3], key[4]))
            else None
        )
        if d and month_start(d) in locked:
            add(Code.E11, "work_date")
        if ex is not None:
            if ex.status == WorkforceStatus.locked:
                add(Code.E11, "work_date")
            elif mode == ImportMode.insert_only:
                add(Code.E10)
            else:
                r["action"] = "replace"
                r["replace_id"] = str(ex.id)
        # warnings
        if hc and mh is not None and mh / hc > s.warn_hours_per_person_day:
            add(Code.W01, "man_hours")
        if eng and site and d and hc is not None:
            hist = [
                h
                for (dd, h) in hc_hist.get((eng.id, site.id), [])
                if d - timedelta(days=7) <= dd < d
            ]
            if hist:
                avg = Decimal(sum(hist)) / Decimal(len(hist))
                if avg > 0 and abs(Decimal(hc) - avg) / avg > Decimal("0.5"):
                    add(Code.W02, "headcount")
        if eng and eng.contractor.status == ContractorStatus.suspended:
            add(Code.W03, "contractor_code")
        if d and (today - d).days > 30:
            add(Code.W04, "work_date")
        if sha_committed:
            add(Code.W05)
        k4 = key[:4]
        shifts = shifts_by_key.get(k4, set()) | file_shifts.get(k4, set())
        if "all" in shifts and ({"day", "night"} & shifts):
            add(Code.W06, "shift")
        r["issues"] = issues
        r["codes"] = sorted(codes)
        errors = any(c.startswith("E") for c in codes)
        r["status"] = (
            ImportRowStatus.error
            if errors
            else ImportRowStatus.warning
            if codes
            else ImportRowStatus.ok
        ).value
        if errors:
            r["action"] = "none"
        r["site_id"] = str(site.id) if site else None
        r["zone_id"] = str(zone.id) if zone else None
        r["engagement_id"] = str(eng.id) if eng else None
    return rows


def _counts(rows: list[dict[str, Any]], inserted: int = 0, replaced: int = 0) -> dict[str, int]:
    return {
        "rows_total": len(rows),
        "rows_ok": sum(1 for r in rows if r["status"] == "ok"),
        "rows_warning": sum(1 for r in rows if r["status"] == "warning"),
        "rows_error": sum(1 for r in rows if r["status"] == "error"),
        "rows_inserted": inserted,
        "rows_replaced": replaced,
    }


def _report(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in rows:
        out.append(
            {
                "row_no": r["row_no"],
                "status": r["status"],
                "codes": r["codes"],
                "issues": r["issues"],
                "work_date": r.get("work_date"),
                "site_code": r.get("site_code"),
                "zone_code": r.get("zone_code"),
                "contractor_code": r.get("contractor_code"),
                "shift": r.get("shift"),
                "action": r.get("action"),
            }
        )
    return out


# ---- read models ---------------------------------------------------------------------------------


def to_read(db: Session, b: WorkforceImportBatch, include_ok: bool = False) -> WorkforceImportRead:
    refs = Refs(db)
    report = [r for r in b.report if include_ok or r["status"] != "ok"]
    report.sort(key=lambda r: ({"error": 0, "warning": 1, "ok": 2}[r["status"]], r["row_no"]))
    return WorkforceImportRead(
        id=b.id,
        project_id=b.project_id,
        file_name=b.file_name,
        file_sha256=b.file_sha256,
        file_size=b.file_size,
        mode=b.mode,
        status=b.status,
        counts=ImportCounts(**b.counts),
        file_issues=[ImportIssue(**i) for i in b.file_issues],
        report=[ImportRowReport(**r) for r in report],
        uploaded_by=refs.user(b.uploaded_by_user_id) or _unknown(),
        created_at=b.created_at,
        expires_at=b.expires_at,
        committed_at=b.committed_at,
    )


def _unknown() -> UserRef:
    return UserRef(id=uuid.UUID(int=0), full_name_en="—")


# ---- API operations ------------------------------------------------------------------------------


def dry_run(
    db: Session,
    p: Principal,
    project_id: uuid.UUID,
    file_name: str,
    content: bytes,
    mode: ImportMode,
) -> WorkforceImportRead:
    project = projects.get_visible(db, p, project_id)
    ensure_open(project)
    p.require(project.id, Capability.workforce_import)
    rows, file_issues = parse(file_name, content)
    sha = hashlib.sha256(content).hexdigest()
    rows = validate(db, p, project, rows, mode, sha)
    b = WorkforceImportBatch(
        project_id=project.id,
        file_name=file_name[:255],
        file_sha256=sha,
        file_size=len(content),
        mode=mode,
        status=ImportStatus.validated,
        counts=_counts(rows),
        file_issues=file_issues,
        report=_report(rows),
        rows=rows,
        uploaded_by_user_id=p.user.id,
        created_at=now(),
        expires_at=now() + EXPIRY,
    )
    db.add(b)
    db.flush()
    audit.record(
        db,
        AuditAction.create,
        p.actor(project.id),
        entity_type=EntityType.workforce_import_batch,
        entity_id=b.id,
        project_id=project.id,
        details={"file_name": b.file_name, "mode": mode.value, "counts": b.counts},
    )
    return to_read(db, b)


def get_batch(db: Session, p: Principal, batch_id: uuid.UUID) -> WorkforceImportBatch:
    b = db.get(WorkforceImportBatch, batch_id)
    if b is None or (
        p.grant(b.project_id, Capability.workforce_import) is None
        and p.grant(b.project_id, Capability.workforce_verify) is None
    ):
        raise deny(
            db,
            p,
            EntityType.workforce_import_batch,
            batch_id,
            b.project_id if b else None,
            "Import batch",
        )
    return b


def get(db: Session, p: Principal, batch_id: uuid.UUID, include_ok: bool) -> WorkforceImportRead:
    return to_read(db, get_batch(db, p, batch_id), include_ok)


def _ensure_validated(db: Session, b: WorkforceImportBatch) -> None:
    if b.status != ImportStatus.validated:
        raise ApiError(
            409,
            ErrorCode.IMPORT_NOT_VALIDATED,
            f"This batch is {b.status.value}; upload the file again.",
            "هذه الدفعة غير قابلة للاعتماد؛ أعد رفع الملف.",
        )
    if now() > b.expires_at:
        b.status = ImportStatus.expired
        db.flush()
        db.commit()  # keep the expiry even though the request fails
        raise ApiError(
            409,
            ErrorCode.IMPORT_EXPIRED,
            "The dry run is older than 60 minutes; upload the file again.",
            "مضى على التحقق أكثر من 60 دقيقة؛ أعد رفع الملف.",
        )


def commit(db: Session, p: Principal, batch_id: uuid.UUID) -> WorkforceImportRead:
    b = get_batch(db, p, batch_id)
    project = db.get(Project, b.project_id)
    assert project is not None  # noqa: S101
    ensure_open(project)
    p.require(project.id, Capability.workforce_import)
    if b.uploaded_by_user_id != p.user.id and not p.is_manager:
        raise forbidden_error("Only the uploader may commit this batch.")
    _ensure_validated(db, b)
    if b.counts.get("rows_error"):
        raise ApiError(
            409,
            ErrorCode.IMPORT_HAS_ERRORS,
            "Fix the errors and upload again.",
            "صحح الأخطاء وأعد الرفع.",
        )
    rows = validate(db, p, project, [dict(r) for r in b.rows], b.mode, b.file_sha256, b.id)
    counts = _counts(rows)
    if counts["rows_error"]:
        b.rows, b.report, b.counts = rows, _report(rows), counts
        db.flush()
        db.commit()
        raise ApiError(
            409,
            ErrorCode.IMPORT_HAS_ERRORS,
            "The data changed since the dry run and new errors appeared.",
            "تغيرت البيانات منذ التحقق وظهرت أخطاء جديدة.",
        )
    inserted = replaced = 0
    for r in rows:
        warnings = [
            {
                "code": i["code"],
                "message": i["message_en"],
                "message_ar": i["message_ar"],
                "field": i["column"],
            }
            for i in r["issues"]
        ]
        values = {
            "site_id": uuid.UUID(r["site_id"]),
            "zone_id": uuid.UUID(r["zone_id"]) if r.get("zone_id") else None,
            "engagement_id": uuid.UUID(r["engagement_id"]),
            "work_date": date.fromisoformat(r["work_date"]),
            "shift": Shift(r["shift"]),
            "no_work": bool(r["no_work"]),
            "headcount": int(r["headcount"]),
            "man_hours": Decimal(r["man_hours"]),
            "toolbox_talks": int(r["toolbox_talks"]),
            "toolbox_attendees": int(r["toolbox_attendees"]),
            "inductions": int(r["inductions"]),
            "training_hours": Decimal(r["training_hours"]),
            "remarks": r.get("remarks"),
        }
        if r.get("action") == "replace":
            ex = db.get(WorkforceReturn, uuid.UUID(r["replace_id"]))
            assert ex is not None  # noqa: S101
            before = wf.snapshot(ex)
            for k, v in values.items():
                setattr(ex, k, v)
            ex.source = WorkforceSource.import_
            ex.import_batch_id = b.id
            ex.status = WorkforceStatus.submitted
            ex.verified_by_user_id = None
            ex.verified_at = None
            ex.warnings = warnings
            replaced += 1
            audit.record(
                db,
                AuditAction.update,
                p.actor(project.id),
                entity_type=EntityType.workforce_return,
                entity_id=ex.id,
                project_id=project.id,
                before=before,
                after=wf.snapshot(ex),
                details={"import_batch_id": str(b.id)},
            )
        else:
            db.add(
                WorkforceReturn(
                    project_id=project.id,
                    source=WorkforceSource.import_,
                    import_batch_id=b.id,
                    status=WorkforceStatus.submitted,
                    warnings=warnings,
                    created_by_user_id=p.user.id,
                    **values,
                )
            )
            inserted += 1
    b.status = ImportStatus.committed
    b.committed_at = now()
    b.rows, b.report = rows, _report(rows)
    b.counts = _counts(rows, inserted, replaced)
    db.flush()
    audit.record(
        db,
        AuditAction.status_change,
        p.actor(project.id),
        entity_type=EntityType.workforce_import_batch,
        entity_id=b.id,
        project_id=project.id,
        details={"committed": True, "counts": b.counts},
    )
    if b.counts["rows_warning"]:
        notify.notify(
            db,
            wf.officers(db, project.id),
            NotificationKind.import_committed_with_warnings,
            f"Workforce import committed with {b.counts['rows_warning']} warning row(s)",
            f"تم اعتماد استيراد القوى العاملة مع {b.counts['rows_warning']} صف بتنبيهات",
            entity_type=EntityType.workforce_import_batch,
            entity_id=b.id,
            project_id=project.id,
        )
    return to_read(db, b)


def discard(db: Session, p: Principal, batch_id: uuid.UUID) -> WorkforceImportRead:
    b = get_batch(db, p, batch_id)
    p.require(b.project_id, Capability.workforce_import)
    if b.status != ImportStatus.validated:
        raise ApiError(
            409,
            ErrorCode.IMPORT_NOT_VALIDATED,
            f"This batch is {b.status.value}.",
            "لا يمكن تجاهل هذه الدفعة.",
        )
    b.status = ImportStatus.discarded
    db.flush()
    audit.record(
        db,
        AuditAction.status_change,
        p.actor(b.project_id),
        entity_type=EntityType.workforce_import_batch,
        entity_id=b.id,
        project_id=b.project_id,
        details={"discarded": True},
    )
    return to_read(db, b)


def list_page(
    db: Session,
    p: Principal,
    project_id: uuid.UUID,
    page: int,
    page_size: int,
    status: ImportStatus | None,
) -> WorkforceImportPage:
    project = projects.get_visible(db, p, project_id)
    g = p.grant(project.id, Capability.workforce_import) or p.grant(
        project.id, Capability.workforce_verify
    )
    if g is None:
        raise forbidden_error()
    B = WorkforceImportBatch  # noqa: N806
    stmt = select(B).where(B.project_id == project.id)
    if g.engagement_ids is not None and not p.grant(project.id, Capability.workforce_verify):
        stmt = stmt.where(B.uploaded_by_user_id == p.user.id)
    if status:
        stmt = stmt.where(B.status == status)
    stmt = stmt.order_by(B.created_at.desc())
    items, total = paginate(db, stmt, page, page_size)
    refs = Refs(db).load(users=[b.uploaded_by_user_id for b in items])
    return WorkforceImportPage(
        items=[
            WorkforceImportSummary(
                id=b.id,
                file_name=b.file_name,
                mode=b.mode,
                status=b.status,
                counts=ImportCounts(**b.counts),
                uploaded_by=refs.user(b.uploaded_by_user_id) or _unknown(),
                created_at=b.created_at,
                committed_at=b.committed_at,
            )
            for b in items
        ],
        total=total,
        page=page,
        page_size=page_size,
    )


def template(fmt_: ExportFormat, headers: str) -> tuple[bytes, str, str]:
    cols = [ar if headers == "ar" else k for k, ar, _ in COLUMNS]
    example = [
        "2026-09-14",
        "ANIA-EXP",
        "S-AIR",
        "",
        "GULFPAVE",
        "night" if headers == "en" else "ليلية",
        600,
        "6000.00",
        "N" if headers == "en" else "لا",
        4,
        560,
        12,
        "48.00",
        "",
    ]
    return export_svc.encode(cols, [example], fmt_, f"workforce-import-template-{headers}")
